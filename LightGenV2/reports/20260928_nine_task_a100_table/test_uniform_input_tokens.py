"""CPU-only validation of the common numerator and scope-preserving workbook edit."""
import json
import math
import os
import hashlib
from pathlib import Path
import unittest
import openpyxl

HERE=Path(__file__).resolve().parent/'frontend_tokens_20261023'


class UniformInputTokensTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        source_path=Path(os.environ.get('UNIFORM_SOURCE_SNAPSHOT',str(HERE/'九任务_核心与含前端_Tokens_20261023.xlsx')))
        cls.old=openpyxl.load_workbook(source_path,data_only=True)
        cls.new=openpyxl.load_workbook(HERE/'九任务_统一输入Tokens_20261024.xlsx',data_only=True)
        cls.meta=json.loads((HERE/'uniform_input_tokens_20261024.json').read_text(encoding='utf-8'))
        if 'UNIFORM_SOURCE_SNAPSHOT' in os.environ:
            if hashlib.sha256(source_path.read_bytes()).hexdigest()!=cls.meta['input_workbook_sha256']:
                raise ValueError('Snapshot is not the exact workbook used by this update')

    def test_costs_and_performance_not_changed(self):
        for old,new in zip(list(self.old)[:2],list(self.new)[:2]):
            for row in range(6,24):
                for col in range(1,13):
                    self.assertEqual(old.cell(row,col).value,new.cell(row,col).value,(new.title,row,col))
                for col in [16,17]:
                    self.assertEqual(old.cell(row,col).value,new.cell(row,col).value)

    def test_same_numerator_for_both_models_in_both_versions(self):
        for sheet in list(self.new)[:2]:
            for i in range(9):
                ours=6+2*i
                self.assertEqual(sheet[f'M{ours}'].value,sheet[f'M{ours+1}'].value)
                self.assertEqual(sheet[f'N{ours}'].value,sheet[f'N{ours+1}'].value)

    def test_token_rate_ratio_equals_latency_speedup(self):
        sheet=self.new.worksheets[0]
        checked=0
        for i in range(9):
            ours,baseline,n=6+2*i,7+2*i,27+i
            if isinstance(sheet[f'C{n}'].value,(int,float)) and isinstance(sheet[f'O{ours}'].value,(int,float)):
                ratio=sheet[f'O{ours}'].value/sheet[f'O{baseline}'].value
                self.assertTrue(math.isclose(ratio,sheet[f'C{n}'].value,rel_tol=1e-12))
                self.assertEqual(sheet[f'H{n}'].value,'一致')
                checked+=1
        self.assertEqual(checked,7)  # 04 scope mismatch and 08 missing token lengths.

    def test_temporal_7136_common_input(self):
        sheet=self.new.worksheets[0]
        self.assertEqual(sheet['M6'].value,7136)
        self.assertEqual(sheet['M7'].value,7136)
        self.assertAlmostEqual(sheet['O7'].value,6004.782217853594)
        self.assertAlmostEqual(sheet['O6'].value,7136/(sheet['H6'].value/1000))

    def test_missing_and_invalid_scope_remain_visible(self):
        core,expanded=list(self.new)[:2]
        self.assertEqual(core['M22'].value,'未保存')
        self.assertEqual(core['M23'].value,'未保存')
        self.assertEqual(core['G31'].value,'范围不匹配')
        for n in range(6,24):
            self.assertEqual(expanded[f'H{n}'].value,'待测')
            self.assertEqual(expanded[f'O{n}'].value,'缺计时或计数')

    def test_native_input_provenance_is_preserved(self):
        for row in range(6,24):
            for col in range(1,11):
                self.assertEqual(self.old['输入依据与Bridge'].cell(row,col).value,self.new['输入依据与Bridge'].cell(row,col).value)

    def test_paired_frontend_evidence_untouched(self):
        old,new=self.old['LGVQ同批边界'],self.new['LGVQ同批边界']
        for row in old:
            for cell in row:
                self.assertEqual(cell.value,new[cell.coordinate].value)


if __name__=='__main__':
    unittest.main()
