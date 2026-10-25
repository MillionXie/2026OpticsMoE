"""Create the Chinese methods appendix from collected, SHA-bound records only."""
from pathlib import Path
from copy import deepcopy
import json
import hashlib
import sys
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE=Path(__file__).resolve().parent
COMPLETE='--complete' in sys.argv
OUT=HERE/'outputs/01a061e2-6f88-7913-87f5-b632000e4c83'/('complete_20261025_01' if COMPLETE else 'collected_20261025_01')
if (OUT/'delivery_manifest.json').exists():raise FileExistsError('Preserve sealed delivery')
DATA=json.loads((OUT/'collected_data_and_audit.json').read_text(encoding='utf8'))
D=Document();section=D.sections[0]
section.page_width,section.page_height=Inches(8.5),Inches(11)
section.top_margin=section.bottom_margin=Inches(.65)
section.left_margin=section.right_margin=Inches(.72)
for name in ('Normal','Title','Heading 1','Heading 2','Caption'):
    s=D.styles[name];s.font.name='Times New Roman';s.font.color.rgb=RGBColor(0,0,0)
    s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
    pp=s.element.find(qn('w:pPr'))
    if pp is not None:
        for e in list(pp.findall(qn('w:pBdr'))):pp.remove(e)
D.styles['Normal'].font.size=Pt(11)
D.styles['Normal'].paragraph_format.line_spacing=1.1
D.styles['Normal'].paragraph_format.space_after=Pt(6)
D.styles['Title'].font.size=Pt(20)
D.styles['Title'].paragraph_format.space_after=Pt(12)
D.styles['Heading 1'].font.size=Pt(14)
D.styles['Heading 1'].paragraph_format.space_before=Pt(10)
D.styles['Heading 1'].paragraph_format.space_after=Pt(7)
D.styles['Heading 2'].font.size=Pt(11.5)
D.styles['Heading 2'].paragraph_format.space_before=Pt(7)
D.styles['Heading 2'].paragraph_format.space_after=Pt(5)
D.core_properties.title='A100时间能耗与比较倍数计算说明'
D.core_properties.author='LightGen2 Research'

def p(text='',style=None):return D.add_paragraph(text,style)
def h(text):return p(text,'Heading 1')
def page(title):D.add_page_break();h(title)
def run_math(text):
    r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=str(text);r.append(t);return r
def element(tag,parts):
    e=OxmlElement(tag)
    for x in parts:e.append(deepcopy(x) if not isinstance(x,str) else run_math(x))
    return e
def sub(base,index):
    e=OxmlElement('m:sSub');e.append(element('m:e',[base]));e.append(element('m:sub',[index]));return e
def sup(base,power):
    e=OxmlElement('m:sSup');e.append(element('m:e',[base]));e.append(element('m:sup',[power]));return e
def frac(num,den):
    e=OxmlElement('m:f');e.append(element('m:num',num));e.append(element('m:den',den));return e
def nary(symbol,lo,hi,body):
    e=OxmlElement('m:nary');pr=OxmlElement('m:naryPr');ch=OxmlElement('m:chr');ch.set(qn('m:val'),symbol);pr.append(ch);e.append(pr)
    e.append(element('m:sub',[lo]));e.append(element('m:sup',[hi]));e.append(element('m:e',body));return e
eq_count=0
def eq(*parts):
    global eq_count
    eq_count+=1;para=p();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before=Pt(4);para.paragraph_format.space_after=Pt(7)
    para.paragraph_format.keep_together=True
    math=element('m:oMath',parts);math.append(run_math(f'   ({eq_count})'))
    mathpara=OxmlElement('m:oMathPara');mathpara.append(math);para._p.append(mathpara)
def table(headers,rows,widths):
    t=D.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    pr=t._tbl.tblPr;borders=OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    pr.append(borders)
    for row_i,values in enumerate([headers]+rows):
        cells=t.rows[0].cells if row_i==0 else t.add_row().cells
        for i,(c,text) in enumerate(zip(cells,values)):
            c.width=Inches(widths[i]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=c._tc.get_or_add_tcPr();margins=OxmlElement('w:tcMar')
            for side in ('top','left','bottom','right'):
                z=OxmlElement('w:'+side);z.set(qn('w:w'),'85');z.set(qn('w:type'),'dxa');margins.append(z)
            tcpr.append(margins);c.text=str(text)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(2);para.paragraph_format.space_before=Pt(2)
                para.alignment=WD_ALIGN_PARAGRAPH.LEFT if i==0 else WD_ALIGN_PARAGRAPH.CENTER
                for r in para.runs:r.font.size=Pt(10);r.font.name='Times New Roman';r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
            fill='233B53' if row_i==0 else ('F0F4F8' if row_i%2==0 else 'FFFFFF')
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),fill);tcpr.append(sh)
            if row_i==0:
                for r in c.paragraphs[0].runs:r.bold=True;r.font.color.rgb=RGBColor(255,255,255)
        if row_i==0:
            repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    p();return t

records={r['key']:r for r in DATA['mainRecords']}
p('A100时间能耗与比较倍数计算说明','Title')
p('完整成本的方法附录 2026年10月25日' if COMPLETE else '已采集数据的方法附录 2026年10月25日')
p('本说明对应两份完整成本Excel及单独的LG系统时间拆分表。九行任务18/18个成本单元均完成，每项3轮各200有效输入；多baseline的36轮和额外80GB Qwen的3轮对照也已完成。LGVQ temporal和spatial在完整输入计费下的时间比分别为0.943与0.543，没有时间加速，能耗比分别为3.521与2.019。' if COMPLETE else '本说明对应两份新建Excel。九行任务已完成14/18个成本单元，其中6行可做Ours与Qwen的成对比较。LGVQ时间质量与空间质量的Ours成本，以及T12双方成本仍待测。多baseline的36轮成本和额外80GB Qwen的3轮对照已完成。')
p('正式性能沿用已确认原代码的完整TEST结果。每轮首200次调用只用于时间、能量和预测一致性核验，不用其子集指标替代完整TEST精度。以下加速与节能倍数均由原始收据计算，没有用历史工作簿的成本补缺。')
table(['任务','时间加速比','能耗降低倍数','能耗降低率','EDP改善比'],[
    [records[id+'|ours']['task'],f"{records[id+'|qwen']['ms']/records[id+'|ours']['ms']:.3f}",f"{records[id+'|qwen']['j']/records[id+'|ours']['j']:.3f}",f"{1-records[id+'|ours']['j']/records[id+'|qwen']['j']:.2%}",f"{records[id+'|qwen']['ms']*records[id+'|qwen']['j']/(records[id+'|ours']['ms']*records[id+'|ours']['j']):.3f}"]
    for id in (('01a','01b','02','03','04','05','06','07','08') if COMPLETE else ('02','03','04','05','06','07'))],[1.8,1.15,1.25,1.2,1.45])
p('表中大于1的时间加速比表示Ours更快，大于1的能耗降低倍数表示Ours每任务能量更低。这些是指定计费边界下的比较：Ours采用实测输入前端和选定神经阶段加参数光学周期的混合硬件投影，baseline采用文件至CPU结果的连续墙钟。两者都不能被改称为同口径整机电表实测。')
p('本文参照ACCEL补充说明4与9的组织方式，分别陈述运算量、耗时分母与各组件能量，再计算吞吐和能效；不移用ACCEL的硬件参数、数值或性能倍数。[1]')

page('1 工作负载与时间边界')
p('每个合格成本单元由3个独立新进程组成，每个进程加载原模型后按固定TEST顺序处理200次调用。没有dummy输入或专用推理预热，首条保留，后续自然变热。模型和权重加载、数据下载、独立离线图库建立，以及结果文件导出不计入在线窗口。ABO文搜图的TEST只有100个不同标题，因此每轮完整遍历两次得到200次计时，正式精度分母仍为100。')
p('baseline窗口从第一条真实文件或UTF8输入开始，到第200条CPU任务结果返回为止。它包含文件读取、解码、采帧、crop/resize、processor/tokenizer、H2D、embedding、原完整网络、任务头、检索/后处理与输出转换。')
eq(sub('T','b,r'),' = ',sub('t','return,r,200'),' − ',sub('t','input,r,1'))
eq(sub('T','b,r'),' = ',nary('∑','i=1','200',[sub('τ','r,i')]),' + ',nary('∑','i=2','200',[sub('g','r,i')]))
p('τ为单次调用区间，g为相邻请求之间的真实间隙。逐条B1路径每轮199个间隙全部保留；LGVQ temporal每轮13整场，对应12个场间隙。没有扣除慢请求或挑选最快一轮。计费采用同一单调时钟，CUDA Event只作配对诊断，神经阶段主计时取同步墙钟。' if COMPLETE else 'τ为单次调用区间，g为相邻请求之间的真实间隙。每轮199个间隙全部保留，不扣除慢请求，不从三轮中挑最快一轮。计费前后使用同一单调时钟。CUDA Event只用于配对诊断，神经阶段的主计时取同步墙钟。')
if COMPLETE:p('数值等价核对与正式计时使用不同进程。T12 LG复用已通过的CPU合同，T12 Qwen与LGVQ在同一实际A100原生精度核对原输入和完整输出。兼容字段cpu_contract只表示路径与SHA，不代表新证明在CPU执行。验证平台变化没有改变模型数学、输入、dtype或正式精度，各计时轮仍独立fresh且零dummy。')
p('Ours的计费电子时间由完整真实输入前端、原选定神经模块、真实merger、CPU输出/检索和对应间隙组成。原FFT数值传播、layout/fusion与被批准单列的并行电子残差仍实际执行并保存raw，但不计入所选混合投影的电子时间。此排除是计费约定，不代表已经实证所有残差都在物理装置中完全重叠。')
eq(sub('T','o,r'),' = ',sub('T','opt,r'),' + ',sub('T','charged,r'))
eq(sub('t','pass'),' = 0.714 + 0.300 + 0.0307 = 1.0447 ms')
p('三个固定分量分别为SLM、曝光和额外CCD时间，ROI为812×812。它们是人工确认的时序参数，不是本轮新测的物理装置端到端墙钟。常规图像查询为6 pass，LSP与SALICON为3 pass，纯标题在线查询为3 pass；标题检索图库的6 pass发生在独立离线阶段。')

page('2 阶段能量与全窗积分')
p('Ours的主值采用批准的阶段参数能量模型。121.746 W为光学器件80.358 W与控制器41.388 W之和；104.230 W为电子计费62.842 W与控制器41.388 W之和。所有时间在能量公式中使用秒。')
eq(sub('E','o,r'),' = 121.746 ',sub('T','opt,r'),' + 104.230 ',sub('T','charged,r'),'   J')
p('该模型不按整个软件仿真窗口持续通电计费，也不把被排除的FFT、layout或并行残差能量加回主值。原软件墙钟及板卡观测作为诊断证据保留，但不是同一物理能量估计的替代项。')
p('baseline的主值为338.2 W主机代理乘同一连续窗口，再加该窗口内实际NVML板卡功率积分。338.2 W不是本轮整机插座电表读数。它与Ours固定阶段功率的估计性质和覆盖范围必须分别说明。')
eq(sub('E','b,r'),' = 338.2 ',sub('T','b,r'),' + ',nary('∫',sub('t','start'),sub('t','end'),[sub('P','board'),'(t) dt']))
eq(sub('E','board,r'),' = ',nary('∑','k','K',[sub('P','k'),' Δ',sub('t','k∩window')]))
p('积分按原审计规则对功率观测作零阶保持，并只累加与真实计费窗口相交的时间。支持年龄上限为250 ms，要求窗口前后有功率观测，不能以无权算术均值替代积分。没有做idle功率扣除。')
eq(sub('P̄','board,r'),' = ',frac([sub('E','board,r')],[sub('T','b,r')]))
table(['量','性质','计入方式'],[
 ['输入前端与选定NN同步墙钟','实测','电子计费时间'],['1.0447 ms/pass','固定参数','按原实际pass或整场次数'],['Ours阶段能量','推导','固定阶段功率乘对应秒数'],['NVML板卡能量','观测积分','baseline同一全窗'],['338.2 W主机能量','代理估计','乘baseline连续秒数']],[2.35,1.2,3.3])

page('3 样本归一化与比较倍数')
p('设每轮有效调用数n=200，重复数R=3。时间和能量先跨三轮求和，再除以600。各轮样本数相等时，这与三轮逐样本均值的平均相同；不是三轮倍数的算术平均。')
eq('t = ',frac([nary('∑','r=1','3',[sub('T','r')])],['600']),' ,   e = ',frac([nary('∑','r=1','3',[sub('E','r')])],['600']))
p('LGVQ每条视频有4个采样帧。每视频成本除以4得到等效单输入帧成本，它是摊销量，不是独立处理一帧任务的实测延迟。纯标题查询的原生单位是每标题，不能说成真实输入图像帧。')
eq(sub('t','equiv-frame'),' = ',frac([sub('t','video')],['4']),' ,   ',sub('e','equiv-frame'),' = ',frac([sub('e','video')],['4']))
p('Temporal Ours必须保留16视频×4帧在整幅相干场联合传播的原合同。200个有效视频需要13个整场，最后一场只有8个有效视频及8个padding槽，仍收取完整场成本，性能与成本的有效视频分母仍为200。')
eq(sub('N','field'),' = ⌈200/16⌉ = 13 ,   ',sub('T','opt,200'),' = 13 × 6 × 1.0447 = 81.4866 ms')
p('上式说明固定执行阶段的归一化规则。完整LGVQ总时间还包含真实前端和选定主体计费时间。不能每视频再收6 pass，也不能将单视频或等效帧吞吐再乘16。' if COMPLETE else '上式只说明固定光学阶段的归一化规则，不是尚未完成的LGVQ正式总时间。不能每视频再收6 pass，也不能将单视频或等效帧吞吐再乘16。')
eq(sub('S','time'),' = ',frac([sub('t','reference')],[sub('t','system')]),' ,   ',sub('S','energy'),' = ',frac([sub('e','reference')],[sub('e','system')]))
eq(sub('η','saving'),' = 1 − ',frac([sub('e','system')],[sub('e','reference')]),' ,   ',sub('S','EDP'),' = ',sub('S','time'),' × ',sub('S','energy'))
p('主表reference为同任务Qwen，system为LG。第二张表同时给出同平台Qwen100参照和LG参照。80GB DS/Kimi相对80GB Qwen100为平台内比较；80GB系统与40GB LG的比值仅作整套配置对照，不把PCIe 40GB/250W与SXM 80GB/400W差异归因于模型。MPS50/100为资源配置，不是目标利用率。' if COMPLETE else '主表reference为同任务Qwen，system为Ours。第二张表Ours成本未完成，因此相关倍数保留待测；已经可计算的比较以同平台Qwen100为reference。80GB DS/Kimi只对80GB Qwen100作平台内比较，不把PCIe 40GB/250W与SXM 80GB/400W差异全部归因于模型。MPS50/100表示执行资源配置，不是目标GPU利用率。')

page('4 共同输入Tokens与算量小计')
p('共同Tokens使用相同首200输入对应的Qwen有效非padding IDs/masks。两模型使用同一分子。LSP/SALICON的共同量是196个视觉patch/图；CLIP和YOLO在第二张表也使用共同Qwen参考分子，不能当作这些模型内部实际token数或语言生成吞吐。')
eq(sub('V','tokens'),' = ',frac([sub('N','common')],[nary('∑','r=1','3',[sub('T','r')])]),' ,   ',frac([sub('V','ours')],[sub('V','qwen')]),' = ',frac([sub('T','qwen')],[sub('T','ours')]))
p('NN运算量按原执行图、实际保存的输入形状、序列长度及分支次数逐项推导。采用1 real MAC=2 OP的约定，不按参数数量估算，不把未执行模块计入。Linear、Conv与attention矩阵分别按对应真实维度计数。')
eq(sub('MAC','Linear'),' = B L ',sub('d','in'),' ',sub('d','out'))
eq(sub('MAC','Conv2d'),' = B ',sub('H','out'),' ',sub('W','out'),' ',sub('C','out'),' (',sub('C','in'),'/g) ',sub('k','h'),' ',sub('k','w'))
eq(sub('MAC','QK+AV'),' = B h ',sub('L','q'),' ',sub('L','k'),' (',sub('d','qk'),' + ',sub('d','v'),')')
p('反卷积遵守原图的输入面积×kernel×输入/输出通道规则，不误用输出面积。DeepSeek按原实际30层与compressed absorbed MLA；Kimi按原FA2的V128到192补齐dense形状计入，并与logical未补齐量分列。')
eq('TOPS = ',frac([sub('OP','matched')],[sup('10','12'),' ',sub('T','matched')]),' ,   TOPS/W = ',frac([sub('OP','matched')],[sup('10','12'),' ',sub('E','matched')]))
p(('18项主表成本完整，其中原14项与T12 LG有NN dense-MAC小计，T12 Qwen和LGVQ LG两项的完整NN小计未核定。T12 LG小计仅本地原图与封存600-shape推导，不冒称已作服务器独立公式回放。' if COMPLETE else '已完成14个主表成本单元均有NN dense-MAC小计。')+'LG实际NN、selected-serial NN、并行残差及其他digital mask contraction分列。小计未覆盖norm、activation、softmax、embedding lookup、插值、索引等全部scalar操作。物理dense-equivalent OP未核定，不以0、容量或GPU FFT替代。GLM/CLIP/YOLO尚无同口径完整NN收据，相应单元标未核定。')
p('协议入口为measurement_protocol_20261024/protocol.json及OpenMoji epoch45修订；完整成本索引为measured_index_complete_cost_18.json与temporal_matrix_complete_cost_08.json。原14成本对象及共同Qwen100锚原样保留。18/18表示时间与能量完整，不等于未知OP已核定。' if COMPLETE else '协议入口为measurement_protocol_20261024/protocol.json及OpenMoji epoch45修订；当前数据事实以partial17、temporal07及其指向的原始收据为准。原协议中的旧排队状态和旧OpenMoji身份由已批准修订覆盖。未完成四成本不能在本说明或工作簿中被称为正式18/18完成。')

page('5 复算实例与原始证据')
o=records['04|ours'];q=records['04|qwen']
p('以ABO文搜图为例。三轮各200次调用，对应同一100标题完整遍历两次。下列数值直接由原三轮收据求和，不挑快轮。')
table(['量','Ours','Qwen'],[
 ['三轮总时间 s',f"{o['seconds']:.9f}",f"{q['seconds']:.9f}"],['三轮总能量 J',f"{o['joules']:.9f}",f"{q['joules']:.9f}"],['每标题 ms',f"{o['ms']:.9f}",f"{q['ms']:.9f}"],['每标题 J',f"{o['j']:.9f}",f"{q['j']:.9f}"],['共同Tokens总数','28,728','28,728']],[2.6,2.1,2.1])
eq(sub('S','time'),' = 37.706426220 / 11.776760241 = 3.201766')
eq(sub('S','energy'),f" = 14.923917347 / 1.282388616 = {q['j']/o['j']:.6f}")
p(f"按未舍入值计算，能耗降低率为{1-o['j']/q['j']:.6%}，EDP改善比为{q['ms']*q['j']/(o['ms']*o['j']):.6f}。倍数与降低率含义不同，11.64倍不等于降低11.64%。")
p((f"本次整理去重{DATA['unique_qualified_points']}个合格测量点，登记{len(DATA['raw_files'])}个原始及证明/收据文件。其中{sum(r['verification']=='封存SHA及字节数已核对' for r in DATA['raw_files'])}个与原封存SHA核对，其余辅助清单或收据登记本轮SHA。" if COMPLETE else '本次整理去重78个合格测量点，共登记915个文件。其中849个与原封存manifest或备份收据的SHA及字节数重新核对；其余66个仅为清单本体、本地独立复算与备份收据，另行登记本轮SHA。')+'没有重新执行已合格神经前向、完整TEST或GPU计时。')
p('检查包括三轮均值、LG阶段能量、baseline主机代理加同窗板卡积分、视频与等效帧四倍关系、两表共享同一40GB Qwen100三轮及共同Tokens率比等于时间反比。Excel中的倍数、百分比和吞吐保留可重算公式，导出缓存另与独立计算交叉检查。')
p('原始文件逐项位置、字节数、SHA256及封存清单来源见两表的原始文件索引页和同目录raw_file_index.csv。原始文件仍位于各任务原run内，没有复制成新的工程或替换旧raw。')
for label,item in [('主表封存索引',DATA['main_index']),('Temporal封存索引',DATA['temporal_index'])]:
    p(label+' '+Path(item['path']).name)
    para=p(item['sha256']);para.paragraph_format.space_after=Pt(5)
    for run in para.runs:run.font.name='Consolas';run.font.size=Pt(9)
p('工作簿用蓝色标示实测分量或原正式指标引用，绿色标示推导，橙色标示固定参数，并保留文字口径说明。未测、未核定与不适用分别显示，均不变成0。OpenMoji采用指定epoch45，仿真0.8765与实验室0.8655分别登记；旧0.888不回填。')
h('参考与复现入口')
p('[1] Chen Y et al. All-analog photoelectronic chip for high-speed vision tasks. Nature 623, 48–57 (2023). DOI 10.1038/s41586-023-06558-8. Supplementary Notes 4 and 9. 本地补充材料SHA256为4f1c5e89c45304e72eb9ac650d9207f36ab3656bd1e4fbde7768bede1795e087。')
if COMPLETE:
    page('6 输入准备对倍率的影响')
    p('LG在本文指整个LightGen系统。当前主表不是仅计GPU操作：它包括真实文件处理、CPU预处理、tokenizer、H2D、embedding、选定GPU模块、CPU输出和请求间隙。同步墙钟包含提交与同步开销，不等于CUDA Event。单独时间表第一列数值保留原主表LG总时间，随后分项可精确加回总数。')
    p('LGVQ temporal每视频404.076793 ms中，视频解码为296.017434 ms。spatial每视频695.970962 ms中，两路原解码合计546.025566 ms。CPU质量特征、processor和tokenizer也计入，不能只看GPU模块的耗时来解释总延迟。')
    eq(sub('S','LG-only'),' = ',frac([sub('T','baseline')],[sub('T','LG'),' − ',sub('T','LG-cut')]))
    eq(sub('S','both-cut'),' = ',frac([sub('T','baseline'),' − ',sub('T','baseline-CPU')],[sub('T','LG'),' − ',sub('T','LG-CPU')]))
    table(['任务','完整输入倍率','只扣LG解码倍率','双方扣CPU倍率'],[['LGVQ temporal','0.943435','3.527878','4.715711'],['LGVQ spatial','0.542559','2.518285','1.530324']],[2.0,1.55,1.55,1.75])
    p('只扣LG解码的倍率保留baseline完整输入耗时，是单侧诊断值，不能当成同口径加速结论。双方扣CPU的情景中，LG扣各frontend_cpu打点段，baseline扣preprocess_start到prepare_end的输入准备段。该剔除只针对已记录区间，余下仍可能有CPU调度、H2D、输出和间隙，不能称为纯GPU模式。')
    p('这两种情景都是原raw的代数剔除，没有重跑准备好输入的推理，缓存、调度或流水线收益未重新测量。正式主表和能耗保留原值；不能直接用缩短的时间乘原平均功率声称新的节能倍数。')
    p('其他任务的input打点常将解码、processor、tokenizer、embedding或H2D合在一起。表中明确写合并段及原stage名称；只有双方CPU边界都独立保存时才给双侧CPU剔除倍率，未独立记录不填0。baseline各配置的CPU输入准备、H2D、完整模型与任务头、检索/后处理及输出和真实间隙均包含在原总时间，模型加载和独立离线图库除外。')

target=OUT/'A100_时间能耗与倍数计算说明_20261025.docx'
D.save(target)
print(json.dumps({'docx':str(target),'equations':eq_count,'explicit_pages':7 if COMPLETE else 6,'status':DATA['status']},ensure_ascii=False))
