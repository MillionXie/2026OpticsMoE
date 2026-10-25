"""Independent file-only interval, power, preserved-output and XLSX checks."""
import csv,json,hashlib,math,sys
from bisect import bisect_right
from pathlib import Path
HERE=Path(__file__).resolve().parent;REPO=HERE.parents[3]
OUT=HERE/'outputs/01a061e2-6f88-7913-87f5-b632000e4c83/first_tf_complete_20261025_01'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def rows(p):
    csv.field_size_limit(10000000)
    with p.open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
def near(a,b):assert math.isclose(a,b,rel_tol=1e-9,abs_tol=1e-6),(a,b)
def point(folder):
    folder=Path(folder);manifest=json.loads((folder/'first_tf_files_manifest.json').read_text())
    for x in manifest:
        p=folder/x['file'];assert p.stat().st_size==x['bytes'] and sha(p)==x['sha256'],str(p)
    observation=json.loads((folder/'first_tf_observer.json').read_text());assert observation['first_block_entries']==200
    assert observation['extra_model_forwards']==observation['explicit_warmup_forwards']==0
    data=rows(folder/'first_tf_calls.csv');power=rows(folder/'telemetry.csv');ts=[float(x['host_monotonic_s']) for x in power]
    duration=board=0.;maxage=0.;previous=-math.inf
    for x in data:
        a=float(x['first_block_start']);b=float(x['cpu_result_return']);assert previous<=a<b;previous=b
        duration+=b-a;near(float(x['core_wall_ms']),(b-a)*1000)
        cursor=a
        while cursor<b:
            i=bisect_right(ts,cursor)-1;assert i>=0 and i+1<len(ts)
            right=min(b,ts[i+1]);assert right>cursor
            board+=float(power[i]['watts'])*(right-cursor);maxage=max(maxage,right-ts[i]);cursor=right
    assert len(data)==200 and maxage<=.25000001
    summary=json.loads((folder/'first_tf_summary.json').read_text());near(duration,summary['wall_s']);near(board,summary['board_j']);near(board+338.2*duration,summary['system_j'])
    import numpy as np
    old=Path(observation['template']);old=REPO/old.as_posix().split('/2026OpticsMoE/',1)[1];old=old.parent
    comparisons=[]
    for name in summary['original_outputs_bitwise_equal']:
        if name.endswith('.npz'):
            with np.load(folder/name) as a,np.load(old/name) as b:
                assert set(a.files)==set(b.files)
                for k in a.files:assert np.array_equal(a[k],b[k]),(name,k)
        else:
            assert [x['ranking'] for x in rows(folder/'calls.csv')]==[x['ranking'] for x in rows(old/'calls.csv')]
        comparisons.append(name)
    report=dict(status='passed',calls=200,wall_s=duration,board_j=board,system_j=board+338.2*duration,
        max_support_age_s=maxage,original_output_comparison=comparisons,no_model_execution=True)
    target=folder/'local_first_tf_independent_audit01.json';assert not target.exists();target.write_text(json.dumps(report,indent=2)+'\n')
    return report
def final():
    from openpyxl import load_workbook
    d=json.loads((OUT/'core_data.json').read_text(encoding='utf8'));assert not any(x.get('pending') for x in d['mainRecords'])
    observed=[];seen=set()
    for r in d['mainRuns']+d['tempRuns']:
        if r['folder'] in seen:continue
        seen.add(r['folder']);folder=REPO/r['folder']
        if r['system']=='ours':
            ledger=rows(folder/'stages.csv');calls=rows(folder/('fields.csv' if (folder/'fields.csv').exists() else 'calls.csv' if (folder/'calls.csv').exists() else 'samples.csv'))
            selected=events=0.
            for i,call in enumerate(calls):
                scope=[x for x in ledger if int(x['sample_index'])==i]
                begin=min(float(x['start_monotonic']) for x in scope if x['stage']==r['start_description'])
                after=[x for x in scope if float(x['start_monotonic'])>=begin]
                previous=begin
                for x in after:
                    a=float(x['start_monotonic']);b=float(x['end_monotonic']);assert a>=previous-1e-9;previous=b
                    assert b<=float(call['cpu_result_return'])
                    near((b-a)*1000,float(x['synchronized_wall_ms']))
                    if x['charged']=='True':
                        selected+=b-a
                        if x['cuda_event_ms']:events+=float(x['cuda_event_ms'])/1000
            near(selected+r['optical_s'],r['wall_s']);near(events+r['optical_s'],r['event_s'])
            near(selected*104.230+r['optical_s']*121.746,r['system_j'])
        else:
            if r.get('new_supplement'):
                audit=json.loads((folder/'local_first_tf_independent_audit01.json').read_text());assert audit['status']=='passed'
                near(audit['wall_s'],r['wall_s']);near(audit['system_j'],r['system_j'])
            else:
                calls=rows(folder/'calls.csv');power=rows(folder/'telemetry.csv');ts=[float(x['host_monotonic_s']) for x in power]
                energy=duration=0.
                for x in calls:
                    a=float(x['first_block_start']);b=float(x['call_return']);duration+=b-a;cursor=a
                    while cursor<b:
                        i=bisect_right(ts,cursor)-1;assert i>=0 and i+1<len(ts);right=min(b,ts[i+1]);assert right>cursor and right-ts[i]<=.25000001
                        energy+=float(power[i]['watts'])*(right-cursor);cursor=right
                near(duration,r['wall_s']);near(338.2*duration+energy,r['system_j'])
        observed.append(r['folder'])
    workbooks=[]
    for name,records,runs,main in [
        ('A100_九任务_首计算模块核心对照_20261025.xlsx',d['mainRecords'],d['mainRuns'],True),
        ('A100_Temporal多基线_首计算模块核心对照_20261025.xlsx',sorted(d['tempRecords'],key=lambda x:x['system']!='ours'),d['tempRuns'],False)]:
        p=OUT/name;w=load_workbook(p,data_only=True);f=load_workbook(p,data_only=False)
        assert not [(s.title,c.coordinate,c.value) for s in w for row in s for c in row if c.data_type=='e']
        for i,r in enumerate(records,6):
            own=[x for x in runs if x['key']==r['key']];s=w['核心计算']
            near(s.cell(i,4).value,sum(x['wall_s'] for x in own));near(s.cell(i,7).value,sum(x['system_j'] for x in own))
            near(s.cell(i,8).value,r['ms']);near(s.cell(i,9).value,r['j'])
        if main:
            s=w['九任务对照']
            for i in range(9):
                a,b=records[2*i:2*i+2];n=i+6;div=4 if i<2 else 1
                near(s.cell(n,6).value,a['ms']/div);near(s.cell(n,7).value,b['ms']/div);near(s.cell(n,8).value,b['ms']/a['ms'])
                near(s.cell(i+19,4).value,b['j']/a['j']);near(s.cell(i+19,6).value,b['j']/a['j']*b['ms']/a['ms'])
        else:
            s=w['Temporal对照'];lg=next(r for r in records if r['system']=='ours')
            for n,r in enumerate(records,6):
                near(s.cell(n,5).value,r['ms']*16);near(s.cell(n,6).value,r['ms']/4);near(s.cell(n,7).value,r['ms']/lg['ms']);near(s.cell(n,9).value,r['j']/lg['j'])
        count=sum(c.data_type=='f' for s in f for row in s for c in row)
        assert all(w[s.title][c.coordinate].value is not None for s in f for row in s for c in row if c.data_type=='f')
        workbooks.append(dict(file=name,bytes=p.stat().st_size,sha256=sha(p),formulas=count))
    report=dict(status='passed',no_model_execution=True,workbooks=workbooks,main_cost_cells=18,
        first_TF_new_baseline_points=21,independently_checked_unique_points=len(observed),all_formula_caches_checked=True,raw_files=len(d['raw_files']))
    (OUT/'independent_delivery_audit.json').write_text(json.dumps(report,indent=2)+'\n');print(json.dumps(report))
if __name__=='__main__':
    if '--point' in sys.argv:print(json.dumps(point(Path(sys.argv[sys.argv.index('--point')+1]))))
    else:final()
