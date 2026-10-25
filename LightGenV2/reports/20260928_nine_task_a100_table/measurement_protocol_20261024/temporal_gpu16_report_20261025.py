"""Read sealed Temporal timing CSVs and audit the new 16-video workbook only."""
import csv
import hashlib
import json
import math
import sys
from collections import defaultdict
from pathlib import Path

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[3]
BASE = HERE / 'outputs/01a061e2-6f88-7913-87f5-b632000e4c83'
SEALED = BASE / 'complete_20261025_01'
OUT = BASE / 'temporal_gpu16_20261025_01'
BOOK = OUT / 'A100_Temporal_16视频_LG显卡时间对完整Qwen_20261025.xlsx'
CORE_OUT = BASE / 'temporal_optical_tf16_20261025_01'
CORE_BOOK = CORE_OUT / 'A100_Temporal_16视频_LG首光模块对Qwen首TF_20261025.xlsx'


def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


def near(a, b):
    assert isinstance(a, (int, float)) and math.isclose(a, b, rel_tol=1e-9, abs_tol=1e-6), (a, b)


def collect():
    source = SEALED / 'collected_data_and_audit.json'
    assert sha(source) == 'f0cd430b99aae15c69ddbeeeee1aa2d2b1334a6e804d7a91097a79c6659d2ffe'
    data = json.loads(source.read_text(encoding='utf-8'))
    catalog = {r['file']: r for r in data['raw_files']}
    used = {}

    def read(folder, name, kind='json'):
        rel = (Path(folder) / name).as_posix()
        p, expected = REPO / rel, catalog[rel]
        assert p.stat().st_size == expected['bytes'] and sha(p) == expected['sha256'], rel
        used[rel] = dict(bytes=p.stat().st_size, sha256=expected['sha256'])
        if kind == 'csv':
            with p.open(encoding='utf-8-sig', newline='') as f:
                return list(csv.DictReader(f))
        return json.loads(p.read_text(encoding='utf-8-sig'))

    lg_runs = [r for r in data['mainRuns'] if r['key'] == '01a|ours']
    q_runs = [r for r in data['tempRuns'] if r['key'] in ('40GB|qwen|100', '40GB|qwen|50')]
    assert len(lg_runs) == 3 and len(q_runs) == 6
    points, stages, reference_ids = [], [], None
    for run in lg_runs:
        folder = run['folder']
        fields, raw = read(folder, 'fields.csv', 'csv'), read(folder, 'stages.csv', 'csv')
        identity, summary = read(folder, 'identity.json'), read(folder, 'summary.json')
        prepared = read(folder, 'prepared_inputs.json')
        ids = [sid for field in fields for sid, valid in zip(json.loads(field['source_ids']), json.loads(field['valid'])) if valid]
        assert reference_ids is None or ids == reference_ids
        reference_ids = ids
        assert len(fields) == 13 and len(ids) == 200
        assert all(json.loads(f['output_shape']) == [1, 16] for f in fields)
        assert all(x['vision_tokens']['shape'] == [1, 16, 4, 49, 1024] for x in prepared)
        assert identity['batch'] == 1
        point = dict(system='LG', cap=100, repeat=run['repeat'], valid=200, fields=13, slots=208,
                     original_ms=run['seconds200'] * 1000, folder=folder, gpu_front_event=0.,
                     gpu_body_event=0., gpu_excluded_event=0., gpu_selected_wall=0.,
                     cpu_prepare=0., cpu_output=0., unspanned_front=0.,
                     gaps=sum(float(f['internal_gap_ms']) for f in fields), fixed=13 * 6 * 1.0447)
        for row in raw:
            wall = float(row['synchronized_wall_ms'])
            near(wall, (float(row['end_monotonic']) - float(row['start_monotonic'])) * 1000)
            event = None if row['cuda_event_ms'] == '' else float(row['cuda_event_ms'])
            charged = row['charged'] == 'True'
            stage = dict(system='LG', repeat=run['repeat'], field=int(row['sample_index']),
                         stage=row['stage'], family=row['family'], region=row['region'], charged=charged,
                         wall_ms=wall, event_ms=event, folder=folder)
            stages.append(stage)
            if event is not None:
                if charged:
                    point['gpu_front_event' if row['region'] == 'frontend' else 'gpu_body_event'] += event
                    point['gpu_selected_wall'] += wall
                else:
                    point['gpu_excluded_event'] += event
            elif charged:
                point['cpu_prepare' if row['region'] == 'frontend' else 'cpu_output'] += wall
        front = sum(float(f['frontend_wall_ms']) for f in fields)
        point['unspanned_front'] = front - sum(s['wall_ms'] for s in stages if s['repeat'] == run['repeat'] and s['region'] == 'frontend')
        assert point['unspanned_front'] >= 0
        assert all(sum(s['family'] == 'optical_simulation_FFT' for s in stages if s['repeat'] == run['repeat'] and s['field'] == i) == 6 for i in range(13))
        near(point['original_ms'], sum(point[k] for k in ('cpu_prepare', 'gpu_selected_wall', 'cpu_output', 'unspanned_front', 'gaps', 'fixed')))
        near(summary['seconds200'] * 1000, point['original_ms'])
        points.append(point)
    for run in q_runs:
        folder = run['folder']
        calls, identity = read(folder, 'calls.csv', 'csv'), read(folder, 'identity.json')
        assert len(calls) == 200 and [c['sample_id'] for c in calls] == reference_ids
        assert identity['batch'] == 1 and all(len(json.loads(c['frame_positions'])) == 4 for c in calls)
        pieces = defaultdict(float)
        for c in calls:
            stamps = [float(c[k]) for k in ('preprocess_start', 'prepare_end', 'model_call_start', 'first_block_start', 'gpu_output', 'call_return')]
            assert all(b >= a for a, b in zip(stamps, stamps[1:]))
            for name, a, b in zip(('cpu_prepare', 'transfer', 'model_front', 'body_head', 'cpu_output'), stamps, stamps[1:]):
                pieces[name] += (b - a) * 1000
        gaps = sum((float(b['preprocess_start']) - float(a['call_return'])) * 1000 for a, b in zip(calls, calls[1:]))
        total = (float(calls[-1]['call_return']) - float(calls[0]['preprocess_start'])) * 1000
        near(total, sum(pieces.values()) + gaps)
        near(total, run['seconds200'] * 1000)
        points.append(dict(system='Qwen', cap=int(run['key'].split('|')[-1]), repeat=run['repeat'],
                           valid=200, fields=200, slots=200, original_ms=total, gaps=gaps, folder=folder, **pieces))
    assert len(points) == 9
    groups = {}
    for s in stages:
        key = s['stage']
        if key not in groups:
            groups[key] = dict(stage=key, family=s['family'], region=s['region'], charged=s['charged'], wall_ms=0., event_ms=0., event_count=0, calls=0)
        g = groups[key]
        assert (g['family'], g['region'], g['charged']) == (s['family'], s['region'], s['charged'])
        g['wall_ms'] += s['wall_ms']; g['calls'] += 1
        if s['event_ms'] is not None:
            g['event_ms'] += s['event_ms']; g['event_count'] += 1
    return dict(points=points, stages=list(groups.values()), source_files=used, no_models_no_GPU=True,
                source_sha256=sha(source), compared_valid_ids=reference_ids,
                scope='Temporal only; LG100 GPU events versus full Qwen40GB MPS100 and50',
                normalization='sum three rounds /600 valid videos *16; retains LG padding cost')


def audit():
    import openpyxl
    assert not (OUT / 'delivery_manifest.json').exists(), 'Preserve sealed report'
    data = collect()
    book = openpyxl.load_workbook(BOOK, data_only=True)
    formulas = openpyxl.load_workbook(BOOK, data_only=False)
    lg = [r for r in data['points'] if r['system'] == 'LG']
    total = lambda key: sum(r[key] for r in lg) * 16 / 600
    selected = total('gpu_front_event') + total('gpu_body_event')
    all_gpu, plus = selected + total('gpu_excluded_event'), selected + total('fixed')
    q = {cap: sum(r['original_ms'] for r in data['points'] if r['system'] == 'Qwen' and r['cap'] == cap) * 16 / 600 for cap in (100, 50)}
    s = book['16视频对照']
    for row, expected in ((6, all_gpu), (7, selected), (8, plus), (9, q[100]), (10, q[50]), (11, total('original_ms'))):
        near(s.cell(row, 2).value, expected)
        near(s.cell(row, 3).value, q[100] / expected)
        near(s.cell(row, 4).value, q[50] / expected)
    near(s['B22'].value, selected)
    near(sum(s.cell(row, 2).value for row in range(16, 22)), selected)
    near(s['B38'].value, total('original_ms'))
    near(sum(s.cell(row, 2).value for row in range(28, 38)), total('original_ms'))
    for cap, column in ((100, 2), (50, 3)):
        near(sum(s.cell(row, column).value for row in range(43, 49)), q[cap])
        near(s.cell(49, column).value, q[cap])
    count = 0
    for sheet in formulas:
        for row in sheet:
            for cell in row:
                cached = book[sheet.title][cell.coordinate]
                assert cached.data_type != 'e', (sheet.title, cell.coordinate)
                if cell.data_type == 'f':
                    count += 1
                    assert cached.value is not None
    report = dict(status='passed', no_models_no_GPU=True, matched_valid_videos_per_configuration=600,
                  LG_fields=39, LG_executed_video_slots=624, Qwen100_calls=600, Qwen50_calls=600,
                  gpu_all_recorded_16_ms=all_gpu, gpu_selected_16_ms=selected,
                  gpu_selected_plus_fixed_16_ms=plus, Qwen_complete16_ms=q, formulas=count,
                  formula_errors=0, original_stage_files_rechecked=True,
                  comparison_is_asymmetric=True, GPU_event_coverage='recorded mutually exclusive scopes, not full hardware kernel trace')
    audit_path = OUT / 'independent_audit.json'
    audit_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    if '--visual-review-accepted' in sys.argv:
        files = [BOOK, OUT / 'analysis_data.json', audit_path]
        manifest = dict(report=report, visual_review_sheets=book.sheetnames,
                        outputs=[dict(path=str(p), bytes=p.stat().st_size, sha256=sha(p)) for p in files])
        (OUT / 'delivery_manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(report))


def collect_core():
    """Re-slice sealed observations, never execute an inference or timing hook.

    LG starts at the first recorded operation in parallel_optics.fields, not
    at an invented Transformer and not at the later FFT propagation boundary.
    Qwen uses the saved native visual-block-0 hook and score-ready endpoint.
    The LG hardware projection retains the existing selected-electronics rule;
    the continuous software interval is retained separately, not concealed.
    """
    import subprocess

    data = collect()
    source = json.loads((SEALED / 'collected_data_and_audit.json').read_text(encoding='utf-8'))
    catalog = {r['file']: r for r in source['raw_files']}
    used = dict(data['source_files'])
    observations, groups, points, evidence = [], {}, [], []

    def read(folder, name, csv_file=False):
        rel = (Path(folder) / name).as_posix()
        p, expected = REPO / rel, catalog[rel]
        assert p.stat().st_size == expected['bytes'] and sha(p) == expected['sha256'], rel
        used[rel] = dict(bytes=expected['bytes'], sha256=expected['sha256'])
        if csv_file:
            with p.open(encoding='utf-8-sig', newline='') as stream:
                return list(csv.DictReader(stream))
        return json.loads(p.read_text(encoding='utf-8-sig'))

    def bound_source(identity, relative, required):
        """Use exact published Git bytes, not Windows checkout CRLF bytes."""
        bound = {k.replace('\\', '/'): v for k, v in identity['sources'].items()}
        assert relative in bound, relative
        blob = subprocess.check_output(['git', 'show', 'HEAD:' + relative], cwd=REPO)
        digest = hashlib.sha256(blob).hexdigest()
        assert digest == bound[relative], (relative, digest, bound[relative])
        text = blob.decode('utf-8-sig')
        assert all(fragment in text for fragment in required), relative
        item = dict(path=relative, bytes=len(blob), sha256=digest,
                    representation='published Git blob, identical to sealed runtime source')
        if item not in evidence:
            evidence.append(item)

    for point in data['points']:
        folder = point['folder']
        identity = read(folder, 'identity.json')
        p = {k: point[k] for k in ('system', 'cap', 'repeat', 'valid', 'fields', 'slots', 'folder')}
        if point['system'] == 'LG':
            bound_source(identity, 'LightGenV2/tasks/t06_video_quality_assessment/models/multivideo9x4.py',
                         ['fields1 = self.parallel_optics.fields(vision)',
                          'encoded = F.softplus(self.width_to_field(tokens.float()))',
                          'self.frame_merger(', 'self.readout('])
            fields = read(folder, 'fields.csv', True)
            stages = read(folder, 'stages.csv', True)
            p.update(event_ms=0., wall_ms=0., fixed_ms=point['fixed'],
                     software_to_cpu_ms=0., cpu_output_ms=0., excluded_event_ms=0.,
                     excluded_wall_ms=0., before_selected_event_ms=0., before_selected_wall_ms=0.)
            for index, field in enumerate(fields):
                body = [r for r in stages if r['region'] == 'body' and int(r['sample_index']) == index]
                starts = [r for r in body if r['stage'] == 'parallel_optics.width_to_field']
                assert len(starts) == 2
                beginning = min(float(r['start_monotonic']) for r in starts)
                after = [r for r in body if float(r['start_monotonic']) >= beginning]
                assert sum(r['family'] == 'optical_simulation_FFT' for r in after) == 6
                assert float(field['frontend_end']) <= beginning < float(field['cpu_result_return'])
                last_end = beginning
                for row in after:
                    start, end = float(row['start_monotonic']), float(row['end_monotonic'])
                    assert beginning <= start < end <= float(field['cpu_result_return'])
                    assert start >= last_end - 1e-9, 'Saved scopes must not overlap'
                    last_end = end
                    wall = float(row['synchronized_wall_ms'])
                    near(wall, (end - start) * 1000)
                    event = float(row['cuda_event_ms']) if row['cuda_event_ms'] else None
                    selected = row['charged'] == 'True'
                    name = row['stage']
                    if name not in groups:
                        groups[name] = dict(stage=name, family=row['family'], charged=selected,
                                            calls=0, event_ms=0., wall_ms=0., event_count=0)
                    g = groups[name]
                    g['calls'] += 1; g['wall_ms'] += wall
                    if event is not None:
                        g['event_ms'] += event; g['event_count'] += 1
                        p['event_ms' if selected else 'excluded_event_ms'] += event
                        p['wall_ms' if selected else 'excluded_wall_ms'] += wall
                    else:
                        assert name == 'output.full_field_CPU_MOS' and selected
                        p['cpu_output_ms'] += wall
                # Earlier body adapters are deliberately outside the new start.
                for row in body:
                    if float(row['start_monotonic']) < beginning and row['charged'] == 'True':
                        assert float(row['end_monotonic']) <= beginning
                        if row['cuda_event_ms']:
                            p['before_selected_event_ms'] += float(row['cuda_event_ms'])
                            p['before_selected_wall_ms'] += float(row['synchronized_wall_ms'])
                software = (float(field['cpu_result_return']) - beginning) * 1000
                p['software_to_cpu_ms'] += software
                observations.append(dict(system='LG', repeat=p['repeat'], field=index,
                    valid=int(sum(json.loads(field['valid']))), slots=16,
                    start_monotonic=beginning, cpu_result_return=float(field['cpu_result_return']),
                    first_optical_stage='parallel_optics.width_to_field', software_to_cpu_ms=software))
            near(p['event_ms'] + p['before_selected_event_ms'], point['gpu_body_event'])
            p['before_selected_event_ms'] += point['gpu_front_event']
            # The complete pre-boundary wall is derived from this point's own
            # source rows, not an equal split of the three different repeats.
            p['before_selected_wall_ms'] += sum(float(r['synchronized_wall_ms']) for r in stages
                if r['region'] == 'frontend' and r['charged'] == 'True' and r['cuda_event_ms'])
            near(p['event_ms'] + p['before_selected_event_ms'], point['gpu_front_event'] + point['gpu_body_event'])
            near(p['wall_ms'] + p['before_selected_wall_ms'], point['gpu_selected_wall'])
        else:
            bound_source(identity, 'LightGenV2/tasks/t06_video_quality_assessment/quality_token_common.py',
                ['model.model.visual.blocks[0].register_forward_pre_hook(self._vision_start)',
                 'self.state["end_event"].synchronize()', 'vision_first_block_to_score_cuda_ms'])
            bound_source(identity, 'LightGenV2/tasks/t06_video_quality_assessment/measure_mps_dataset_once.py',
                ["timer.state['vision_host_started'] if args.model=='qwen'", 'timing'])
            calls = read(folder, 'calls.csv', True)
            p.update(event_ms=0., wall_ms=0., fixed_ms=None,
                     software_to_cpu_ms=0., cpu_output_ms=0., excluded_event_ms=None,
                     excluded_wall_ms=None, before_selected_event_ms=None, before_selected_wall_ms=None)
            for call in calls:
                start, end, returned = [float(call[key]) for key in ('first_block_start', 'gpu_output', 'call_return')]
                event = float(call['diagnostic_first_block_cuda_ms'])
                assert start < end <= returned and event > 0
                near((end - start) * 1000, float(call['legacy_host_ms']))
                p['event_ms'] += event; p['wall_ms'] += (end - start) * 1000
                p['software_to_cpu_ms'] += (returned - start) * 1000
                p['cpu_output_ms'] += (returned - end) * 1000
                observations.append(dict(system='Qwen', cap=p['cap'], repeat=p['repeat'],
                    sample_id=call['sample_id'], first_block_start=start, gpu_output=end,
                    call_return=returned, cuda_event_ms=event))
            near(p['wall_ms'], point['body_head'])
            near(p['cpu_output_ms'], point['cpu_output'])
        points.append(p)
    assert len(observations) == 39 + 600 + 600
    return dict(points=points, stages=list(groups.values()), observations=observations,
        source_files=used, source_code_evidence=evidence,
        source_sha256=data['source_sha256'], no_models_no_GPU=True,
        matched_valid_videos_per_configuration=600, lg_fields=39, lg_executed_video_slots=624,
        start_LG='first parallel_optics.width_to_field in each unchanged 16video field, including optical-input mapping',
        start_Qwen='native first visual Transformer block model.model.visual.blocks[0]',
        LG_rule='existing selected electronic scopes after start + six fixed optical passes; not continuous software E2E',
        normalization='sum three repeats /600 valid videos *16; equivalent valid frame divides by64; retain all padding',
        fixed_optical_ms_per_field=6 * 1.0447)


def audit_core():
    import openpyxl
    assert not (CORE_OUT / 'delivery_manifest.json').exists(), 'Preserve sealed report'
    data = collect_core()
    book = openpyxl.load_workbook(CORE_BOOK, data_only=True)
    formulas = openpyxl.load_workbook(CORE_BOOK, data_only=False)
    s = book['16视频核心对照']
    lg = [p for p in data['points'] if p['system'] == 'LG']
    sums = lambda points, key: sum(p[key] for p in points) * 16 / 600
    lg_event = sums(lg, 'event_ms') + sums(lg, 'fixed_ms')
    lg_wall = sums(lg, 'wall_ms') + sums(lg, 'fixed_ms')
    expected = []
    for row, system, cap in ((6, 'LG', 100), (7, 'Qwen', 100), (8, 'Qwen', 50)):
        points = [p for p in data['points'] if p['system'] == system and p['cap'] == cap]
        event = lg_event if system == 'LG' else sums(points, 'event_ms')
        wall = lg_wall if system == 'LG' else sums(points, 'wall_ms')
        for rr, value, base in ((row, event, lg_event), (row + 9, wall, lg_wall)):
            near(s.cell(rr, 4).value, value)
            near(s.cell(rr, 5).value, value / 64)
            near(s.cell(rr, 6).value, value / base)
        expected.append(dict(system=system, cap=cap, event16_ms=event, wall16_ms=wall,
                             event_relative_to_LG=event / lg_event, wall_relative_to_LG=wall / lg_wall))
    near(s['B33'].value, sums(lg, 'wall_ms'))
    near(s['C33'].value, sums(lg, 'event_ms'))
    near(s['B34'].value, sums(lg, 'fixed_ms'))
    near(s['C34'].value, sums(lg, 'fixed_ms'))
    near(s['B35'].value, lg_wall); near(s['C35'].value, lg_event)
    stage_names = ('parallel_optics.width_to_field', 'parallel_optics.tokens_to_field',
                   'parallel_optics.expert_readout', 'parallel_optics.global_readout', 'frame_merger',
                   'serial_optics.width_to_field', 'serial_router.width_to_field',
                   'serial_router.tokens_to_field', 'readout')
    for row, name in enumerate(stage_names, 24):
        stage = next(item for item in data['stages'] if item['stage'] == name)
        near(s.cell(row, 2).value, stage['wall_ms'] * 16 / 600)
        near(s.cell(row, 3).value, stage['event_ms'] * 16 / 600)
        near(s.cell(row, 4).value, stage['event_ms'] * 16 / 600 / lg_event)
    near(s['D33'].value, sums(lg, 'event_ms') / lg_event)
    near(s['D34'].value, sums(lg, 'fixed_ms') / lg_event)
    near(s['D35'].value, 1.)
    near(s['B40'].value, sums(lg, 'software_to_cpu_ms'))
    for cap, column in ((100, 3), (50, 4)):
        points = [p for p in data['points'] if p['system'] == 'Qwen' and p['cap'] == cap]
        near(s.cell(40, column).value, sums(points, 'software_to_cpu_ms'))
    count = 0
    for sheet in formulas:
        for row in sheet:
            for cell in row:
                cached = book[sheet.title][cell.coordinate]
                assert cached.data_type != 'e', (sheet.title, cell.coordinate)
                if cell.data_type == 'f':
                    count += 1
                    assert cached.value is not None
    # Old workbook and its source observations remain immutable.
    old = OUT / 'delivery_manifest.json'
    assert sha(old) == 'ffc87fbd4aef5e0364429ef50eb9f5754f1ab814e37ba8d24145acd324358eb2'
    old_manifest = json.loads(old.read_text(encoding='utf-8'))
    for item in old_manifest['outputs']:
        p = Path(item['path'])
        assert p.stat().st_size == item['bytes'] and sha(p) == item['sha256']
    report = dict(status='passed', no_models_no_GPU=True, first_sample_and_padding_retained=True,
        matched_valid_videos_per_configuration=600, actual_LG_fields=39, actual_LG_video_slots=624,
        actual_Qwen_calls_per_configuration=600, raw_observations_checked=len(data['observations']),
        native_first_visual_TF_source_bound=True, LG_first_optical_input_mapping_source_bound=True,
        old_workbook_preserved=True, formulas=count, formula_errors=0, results=expected,
        LG_metric='existing segmented hardware projection, not continuous physical end-to-end measurement')
    audit_path = CORE_OUT / 'independent_audit.json'
    audit_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    if '--visual-review-accepted' in sys.argv:
        files = [CORE_BOOK, CORE_OUT / 'analysis_data.json', audit_path]
        manifest = dict(report=report, visual_review_sheets=book.sheetnames,
                        outputs=[dict(path=str(p), bytes=p.stat().st_size, sha256=sha(p)) for p in files])
        (CORE_OUT / 'delivery_manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(report))


if __name__ == '__main__':
    if '--core-data' in sys.argv:
        print(json.dumps(collect_core(), ensure_ascii=True))
    elif '--core-audit' in sys.argv:
        audit_core()
    elif '--data' in sys.argv:
        print(json.dumps(collect(), ensure_ascii=True))
    else:
        audit()
