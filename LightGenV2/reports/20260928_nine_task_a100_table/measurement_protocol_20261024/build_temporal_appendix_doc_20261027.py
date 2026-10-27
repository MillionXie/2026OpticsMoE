"""Create and independently inspect the timing appendix; no model/GPU access."""
import argparse
import hashlib
import json
import math
from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches,Pt,RGBColor
from PIL import Image,ImageDraw,ImageFont

NAME='Temporal_LightGen_附录计时与统计方法_20261027.docx'
ACCEL_URL='https://pmc.ncbi.nlm.nih.gov/articles/PMC10620079/'


def digest(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def math_run(s):
    r=OxmlElement('m:r')
    if s=='max' or str(s).endswith(' ms'):
        pr=OxmlElement('m:rPr');style=OxmlElement('m:sty');style.set(qn('m:val'),'p');pr.append(style);r.append(pr)
    t=OxmlElement('m:t');t.text=str(s);r.append(t);return r


def sub(base,index):
    n=OxmlElement('m:sSub');e=OxmlElement('m:e');e.append(math_run(base));i=OxmlElement('m:sub');i.append(math_run(index));n.extend([e,i]);return n


def fraction(top,bottom):
    f=OxmlElement('m:f');n=OxmlElement('m:num');d=OxmlElement('m:den')
    for x in top:n.append(x)
    for x in bottom:d.append(x)
    f.extend([n,d]);return f


def summation(index,upper,body):
    n=OxmlElement('m:nary');pr=OxmlElement('m:naryPr');ch=OxmlElement('m:chr');ch.set(qn('m:val'),'∑');pr.append(ch)
    loc=OxmlElement('m:limLoc');loc.set(qn('m:val'),'undOvr');pr.append(loc);n.append(pr)
    lo=OxmlElement('m:sub');lo.append(math_run(index));hi=OxmlElement('m:sup');hi.append(math_run(upper));e=OxmlElement('m:e')
    for x in body:e.append(x)
    n.extend([lo,hi,e]);return n


def equation(doc,items,number):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(9)
    om=OxmlElement('m:oMathPara');m=OxmlElement('m:oMath')
    for x in items:m.append(x)
    m.append(math_run('    ('+str(number)+')'));om.append(m);p._p.append(om)


def paragraph(doc,text,style=None):
    p=doc.add_paragraph(text,style);p.paragraph_format.widow_control=True;return p


def table(doc,headers,rows,widths,numeric=()):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
    t._tbl.tblPr.append(borders)
    margin=OxmlElement('w:tblCellMar')
    for edge in ['top','left','bottom','right']:
        v=OxmlElement('w:'+edge);v.set(qn('w:w'),'75');v.set(qn('w:type'),'dxa');margin.append(v)
    t._tbl.tblPr.append(margin)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for i,values in enumerate([headers]+rows):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        for j,(cell,value,width) in enumerate(zip(cells,values,widths)):
            cell.width=Inches(width);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text=str(value);p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=Pt(13)
            snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');p._p.get_or_add_pPr().append(snap)
            p.alignment=WD_ALIGN_PARAGRAPH.RIGHT if j in numeric and i else WD_ALIGN_PARAGRAPH.CENTER if i==0 else WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.size=Pt(10);run.font.bold=i==0
                if i==0:run.font.color.rgb=RGBColor(255,255,255)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'243B53' if i==0 else 'F1F5F8' if i%2 else 'FFFFFF');cell._tc.get_or_add_tcPr().append(shade)
    return t


def flow_figure(path):
    im=Image.new('RGB',(1600,610),'white');d=ImageDraw.Draw(im)
    font=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',30)
    small=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',27)
    def box(bounds,text):
        d.rounded_rectangle(bounds,radius=8,fill='#F1F5F8',outline='#425D73',width=2)
        x0,y0,x1,y1=bounds;b=d.multiline_textbbox((0,0),text,font=font,spacing=5,align='center')
        d.multiline_text(((x0+x1-b[2])/2,(y0+y1-b[3])/2-3),text,font=font,fill='#142C3E',align='center',spacing=5)
    def arrow(a,b):
        d.line((a,b),fill='#425D73',width=3)
        x,y=b;dx,dy=b[0]-a[0],b[1]-a[1];norm=math.hypot(dx,dy);ux,uy=dx/norm,dy/norm
        d.polygon([(x,y),(x-13*ux+7*uy,y-13*uy-7*ux),(x-13*ux-7*uy,y-13*uy+7*ux)],fill='#425D73')
    box((35,15,460,90),'文本分词（Tokenizer）')
    box((35,145,460,225),'文本Embedding')
    box((580,145,1030,225),'视觉Embedding')
    arrow((247,90),(247,145))
    d.multiline_text((1090,140),'两路Embedding可重叠\n各分支完成后汇合',font=small,fill='#425D73',spacing=10)
    box((110,295,1025,390),'视觉Adapter\n前端位置处理、特征转换、视觉调制及光场输入映射')
    arrow((247,225),(247,295));arrow((805,225),(805,295))
    box((20,475,290,575),'视觉光电处理')
    box((350,475,710,575),'Merger与\n多模态序列准备')
    box((770,475,1060,575),'语言Adapter')
    box((1120,475,1580,575),'语言光电处理\n及Decoder与最终评分')
    arrow((155,390),(155,475));arrow((290,525),(350,525));arrow((710,525),(770,525));arrow((1060,525),(1120,525))
    im.save(path)


def build(folder,revise=False):
    folder=Path(folder).resolve();target=folder/NAME
    if (folder/'delivery_manifest.json').exists():raise ValueError('Never modify a sealed delivery')
    if target.exists() and not revise:raise FileExistsError(target)
    data=json.loads((folder/'static_values.json').read_text(encoding='utf8'));p=data['appendix_table'];total=data['rows'][0][5]
    if data.get('edit_mode')=='raw_appendix':
        return edit_raw_appendix(folder,data,target)
    if data.get('edit_mode')=='optical_control_to_electronics':
        return edit_optical_control(folder,data,target)
    if len(p['rows'])!=18 or not p['rows'][0][1].endswith('视觉Adapter'):raise ValueError('Functional adapter grouping required')
    if not math.isclose(math.fsum(r[2] for r in p['rows'][:-1]),total,abs_tol=1e-11):raise ValueError('Timing table does not close')
    image=folder/'timing_workflow.png';flow_figure(image)
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.7)
    for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
        style=doc.styles[name];style.font.name='Microsoft YaHei';style.font.color.rgb=RGBColor(0,0,0)
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
        style.font.size=Pt(11 if name=='Normal' else 19 if name=='Title' else 14 if name=='Heading 1' else 11)
        style.paragraph_format.space_after=Pt(5);style.paragraph_format.line_spacing=Pt(15.5)
        style.paragraph_format.space_before=Pt(0 if name in ['Normal','Title','Caption'] else 10)
        style.font.bold=name in ['Heading 1','Heading 2']
        snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');style._element.get_or_add_pPr().append(snap)
    doc.styles['Caption'].font.size=Pt(10)
    for border in list(doc.styles._element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    doc.core_properties.title='LightGen时间统计与并行计算方法'
    doc.core_properties.subject='每16视频64帧的等效时间分解与原始三轮核算'
    doc.core_properties.author=''
    doc.add_paragraph('LightGen时间统计与并行计算方法','Title')
    paragraph(doc,f'本附录定义LightGen各步骤的功能边界、实际先后关系和并行计时方法。三轮记录得到的平均等效时间为{total:.6f} ms/16视频（64帧）。各步骤均从同一组运行记录核算，分块相加与总时间一致。')
    doc.add_heading('测量工作量与报告单位',level=1)
    paragraph(doc,'每组包含16个视频，每个视频取4帧，共64帧。每轮对512个原TEST视频循环4次，形成2048次视频调用、128组。共进行3轮完整计时，总计6144次视频调用、384组。循环用于计时，不代表新增独立测试样本。')
    paragraph(doc,'“等效时间”表示整轮流水运行时间折算到每组16视频后的平均值。它包含原计时窗口内首次运行和最后返回的影响，不等于首组视频从输入到返回的单独延迟。电子阶段记录来自A100 40GB PCIe，原功率上限250 W、MPS 100%。')
    doc.add_heading('功能模块与实际依赖顺序',level=1)
    doc.add_picture(str(image),width=Inches(7.05))
    # Inline drawings need an automatic line height, not the exact prose height.
    doc.paragraphs[-1].paragraph_format.line_spacing=1.0
    paragraph(doc,'图1 本场数据依赖示意。Adapter统一指输入光场前的适配功能，不照搬程序类名。调度与同步贯穿整轮，图中不画成Decoder之后的串行尾段。','Caption')
    paragraph(doc,'文本分词先生成token ID。文本Embedding与视觉Embedding可并行；随后视觉Adapter完成前端位置处理、特征转换、缩放和偏置、视觉特征调制及首次视觉光场输入映射。该功能块包含源码中用于处理文本特征的相关操作，功能边界与表中的视觉Adapter一致。')
    paragraph(doc,'视觉光电处理完成后，Merger聚合视觉特征，并与已经得到的文本特征组成多模态序列。语言Adapter位于Merger之后、语言Router光传播之前，负责相应光场输入适配；它不是再次分词或再次进行文本Embedding。后续周期的输入适配保留在相应OEO电子处理边界内。')

    doc.add_page_break();doc.add_heading('并行后的等效时间计算',level=1)
    paragraph(doc,'我们将每轮记录表示为依赖图。一个节点对应一次计算或必要控制动作，边表示输入依赖、同一执行流的先后关系或槽位复用约束。每个节点只能在所有前驱完成后开始；并行支路汇合时取较晚完成的时间，而不是把支路耗时相加。')
    paragraph(doc,'记第r轮节点v的核算时长为dᵣ(v)，其前驱集合为pred(v)，最早完成时刻为Eᵣ(v)。无前驱节点的前驱完成时刻取0。整轮时间Lᵣ由最后完成的节点确定：')
    equation(doc,[sub('E','r'),math_run('(v) = '),sub('d','r'),math_run('(v) + '),sub('max','u∈pred(v)'),math_run(' '),sub('E','r'),math_run('(u),    '),sub('L','r'),math_run(' = '),sub('max','v'),math_run(' '),sub('E','r'),math_run('(v)')],1)
    paragraph(doc,'决定整轮完成时间的最长依赖链称为关键路径。采用与原记录一致的同一条关键路径Cᵣ进行分块，每个保留节点只归属一个步骤。令Vₖ为归属步骤k的节点集合，步骤k的等效时间定义为：')
    equation(doc,[sub('t','eq,k'),math_run(' = '),fraction([summation('r=1','3',[summation('v∈Cᵣ∩Vₖ','',[sub('d','r'),math_run('(v)')])])],[math_run('384')])],2)
    paragraph(doc,'这一定义是总运行时间的互斥分解，不是某步骤单独运行的耗时，也不能解释为删除该步骤就必然减少同样的时间。CPU提交、必要主机间隙和就绪排队若位于所采用的依赖链上，已计入对应步骤。剩余全程控制只计入“全程调度与同步”。')
    equation(doc,[sub('T','16'),math_run(' = '),fraction([summation('r=1','3',[sub('L','r')])],[math_run('3×128')]),math_run(' = '),summation('k','',[sub('t','eq,k')])],3)
    paragraph(doc,'表1 三轮完整运行时间及每组平均值。起点均为原计时窗口t=0，未删除首组或选择较快窗口。','Caption')
    rounds=data['phase_table']['rounds']
    table(doc,['轮次','组数','整轮时间/ms','等效时间/ms/16视频'],[[str(int(x[0])),str(int(x[1])),f'{x[3]:.9f}',f'{x[4]:.9f}'] for x in rounds],[.6,.6,2.2,3.7],(1,2,3))
    equation(doc,[sub('T','16'),math_run(' = '),fraction([math_run(f'{math.fsum(x[3] for x in rounds):.12f}')],[math_run('384')]),math_run(f' = {total:.12f} ms')],4)
    paragraph(doc,'所有计算先使用全精度，再统一格式化显示。表中6位或9位小数的末位差异仅来自四舍五入，不对任何原始事件作补差、平移时钟或人为扣减。')

    doc.add_page_break();doc.add_heading('LightGen分阶段等效时间',level=1)
    paragraph(doc,'表2 每组16视频、每视频4帧，共64帧。下列17个联合处理或控制项的等效时间可直接相加；总计行不再重复相加。预处理内部步骤未被原始事件完全分离，因此保留联合计时。','Caption')
    headers=['模块','处理步骤','等效时间/ms\n每16视频64帧']
    timing_rows=[[r[0],r[1],f'{r[2]:.6f}'] for r in p['rows']]
    table(doc,headers,timing_rows,[1.2,4.3,1.6],(2,))
    paragraph(doc,'预处理的联合等效时间为9.751447 ms，包含文本分词、两路Embedding及视觉Adapter。没有按参数量、FLOPs或人为比例拆出未独立测量的子项。语言Adapter只列原有独立事件边界内、进入语言Router前的适配。')
    paragraph(doc,'OEO指光—电—光衔接所需的电子处理。表内沿用原事件边界，包含必要的读出、路由、融合和下段输入准备，不等同于单个ADC或DAC器件的延时。Merger行亦为联合事件范围，包含序列位置、掩码及同一边界内必要的输入准备。')

    doc.add_page_break();doc.add_heading('并行关系与计时范围',level=1)
    paragraph(doc,'同场依赖。视觉与语言各包含Router、Expert、Global三个光周期；必要OEO电子处理保留在对应光周期之后。Merger依赖视觉处理完成；语言Adapter依赖Merger产物；Decoder与最终评分依赖语言处理完成。没有将这些有输入依赖的步骤假定为同时完成。')
    paragraph(doc,'跨场流水。后一场的文本分词及前处理可以与前一场主体计算重叠。视觉Embedding中的原完整patch分支保留实际fork/join及拼接成本，文本Embedding与视觉Embedding之间保留原汇合约束。槽位复用仍须等待该槽前次评分完成。式（1）的依赖关系处理这些重叠，不再从总时间第二次扣除。')
    paragraph(doc,'全程控制。主机发出计算指令所需的CPU提交、动作间隙及工作就绪后的排队，按原目标事件归入相应步骤。不能归入单一步骤的整轮进入、场间主机间隙、槽位等待、捕获周边间隙以及同步返回，合为0.737746 ms的“全程调度与同步”。该行是全程汇总，不是最后额外执行的一段，也不重复包含已归入各步骤的控制开销。')
    doc.add_heading('电子阶段与光周期的核算边界',level=1)
    paragraph(doc,'电子阶段沿用三轮A100原始事件。每个光周期采用既定器件时序：1000/1400 ms、0.300000 ms和0.030700 ms之和，即1.044985714 ms。六个周期均保留整幅光场传播及原依赖。此数值由电子阶段计时与固定光周期组合核算，不是物理光电整机的直接端到端实测。')
    paragraph(doc,'计时起点至最终CPU评分返回的原边界保持不变。模型加载、原图准备和批量H2D在窗口前。首三组直接执行以及两次真实CUDA图捕获对应的输出仍计入工作量，没有额外预热。仅按既定口径处理指定FFT仿真、对应提交、ROI辅助步骤、电子残差及原H2D节点，并排除已批准的一次性电子图构建函数；其相邻观测间隙及必要OEO电子处理没有一并删除。本次整理未新增排除项。')
    doc.add_heading('总时间核对与资料来源',level=1)
    equation(doc,[sub('T','16'),math_run(' = '),sub('t','预处理'),math_run(' + '),sub('t','视觉'),math_run(' + '),sub('t','融合'),math_run(' + '),sub('t','语言'),math_run(' + '),sub('t','解码'),math_run(' + '),sub('t','调度')],5)
    paragraph(doc,'上式依次对应预处理、视觉光电处理、多模态融合、语言光电处理、Decoder与最终评分、全程调度与同步。六模块全精度相加等于14.735726904997131 ms，亦等于三轮整轮时间之和除以384。Excel保存同一套固定数值，不依赖其他工作表的公式。')
    paragraph(doc,'内部核对记录对应A100运行a100_lgvq_async_20261026_17、18、19，各轮原始事件、逐字节输出检查、独立审计及v3时间账本均保留。电子阶段、计时边界和输入顺序以这些运行记录为依据。')
    paragraph(doc,'分段计时与并行覆盖的报告方式参考ACCEL的系统时序说明：被另一阶段覆盖的时间不再重复累计。本附录没有套用ACCEL的器件参数或延时数值。[1]')
    paragraph(doc,'[1] Chen Y et al. All-analog photoelectronic chip for high-speed vision tasks. Nature 623, 48–57 (2023). doi:10.1038/s41586-023-06558-8. 系统时序说明见Computing speed and efficiency及Fig. 6。')
    paragraph(doc,ACCEL_URL)
    for border in list(doc._element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    doc.save(target)
    print(json.dumps(dict(status='created',path=str(target),equations=5,timing_rows=len(timing_rows)),ensure_ascii=True))


def replace_plain_text(paragraph,text):
    """Retain the paragraph and its first run formatting during a narrow edit."""
    if paragraph.runs:
        paragraph.runs[0].text=text
        for run in paragraph.runs[1:]:run.text=''
    else:paragraph.add_run(text)


def edit_optical_control(folder,data,target):
    source=Path(data['base_word_document'])
    if digest(source)!=data['base_word_sha256']:raise ValueError('Source Word changed')
    doc=Document(source);p=data['appendix_table']
    if len(doc.tables)!=2 or len(doc.tables[1].rows)!=19:raise ValueError('Unexpected source Word structure')
    for row,expected in zip(doc.tables[1].rows[1:],p['rows']):
        for cell,value in zip(row.cells,[expected[0],expected[1],f'{expected[2]:.6f}']):
            if cell.text!=value:
                if len(cell.paragraphs)!=1:raise ValueError('Unexpected cell paragraphs')
                replace_plain_text(cell.paragraphs[0],value)
    changed=[]
    for para in doc.paragraphs:
        if para.text.startswith('OEO指光—电—光衔接所需的电子处理。'):
            replace_plain_text(para,para.text+' 光传播行仅列纯光学等效部分；原归入该光传播行的CPU提交、主机间隙和就绪排队，统一合入紧接其后的同名OEO电子处理行。每个光、电两行小计与总时间均保持不变。')
            changed.append('oeo_scope')
        elif para.text.startswith('电子阶段沿用三轮A100原始事件。每个光周期'):
            replace_plain_text(para,para.text+' 表2六个光传播行均为0.348329 ms/16视频场（64帧）的纯光学等效贡献；它是并行核算后的整轮平均分摊，不是单次光周期的1.044985714 ms延时。')
            changed.append('optical_units')
    if changed!=['oeo_scope','optical_units']:raise ValueError('Expected two source paragraphs')
    doc.save(target)
    print(json.dumps(dict(status='edited',path=str(target),source_sha256=digest(source),
        equations=5,timing_rows=len(p['rows']),preserved_original_document=True),ensure_ascii=True))


def edit_raw_appendix(folder,data,target):
    source=Path(data['base_word_document'])
    if digest(source)!=data['base_word_sha256']:raise ValueError('Source Word changed')
    doc=Document(source);raw=data['raw_appendix'];p=data['appendix_table']
    for row,expected in zip(doc.tables[1].rows[1:],p['rows']):
        if [c.text for c in row.cells]!=[expected[0],expected[1],f'{expected[2]:.6f}']:raise ValueError('Existing Word and edited Excel differ')
    found=0
    for para in doc.paragraphs:
        if para.text.startswith('电子阶段沿用三轮A100原始事件。每个光周期'):
            replace_plain_text(para,para.text.replace('1000/1400 ms、0.300000 ms和0.030700 ms',
                'SLM刷新1000/1400 ms、CCD曝光0.300000 ms和CCD读出0.030700 ms'))
            found+=1
    if found!=1:raise ValueError('Original optical boundary paragraph missing')
    doc.add_page_break();doc.add_heading('原始分块与等效时间核对',level=1)
    paragraph(doc,'表3将三轮原始计算事件按表2相同功能边界合并。下列各值均除以384场，单位为ms/16视频场（64帧）。控制开销单独列出，光传播使用固定器件时序。原计算累计未处理并行覆盖，不能直接相加作为总时间。')
    paragraph(doc,'对于每一行，等效时间等于并行后计算与并行后控制之和。三轮各自保留128场的逐轮分块数据，汇总时按合计时间除以合计场数核算，不挑选较快窗口。')
    rows=[[r[0]+' '+r[1].replace('\n',' '),*[f'{r[j]/384:.6f}' for j in [2,3,4]],f'{r[6]:.6f}'] for r in raw['summary_rows']]
    table(doc,['模块与步骤','原计算累计\n未去重','并行后\n计算','并行后\n控制','等效时间'],rows,[3.1,1.,1.,1.,1.],(1,2,3,4))
    paragraph(doc,'“原计算累计”是原电子事件或固定光周期的时长之和，再除以384；“并行后计算”和“并行后控制”只计原依赖链中决定完成时间的部分。两列逐行相加得到表2等效时间，所有行合计14.735726904997131 ms。')
    paragraph(doc,'预处理保留联合计时，不将未独立测量的Embedding或Adapter子操作人为分配时间。光传播行的提交与排队归入对应OEO电子处理，全程调度不再次包含这些部件内控制。')
    doc.add_page_break();doc.add_heading('全程调度原始记录与汇总核对',level=1)
    paragraph(doc,f'原始调度数据共{raw["original_control_records"]}条。Excel的“调度原始明细”每行保存一个原始同步调用或相邻主机间隙，保留轮次、从0计的场次、原始起止时间、原始时长、原始节点ID和依赖节点。未按场平均或合并原始记录。所有起止时间相对该轮原始host_start。')
    paragraph(doc,'同步调用的原始区间可能包含等待GPU完成的阻塞。该等待由原依赖边表达。依赖分离后的节点时长仅保留GPU就绪后仍必要的返回开销，随后再按原并行依赖链核算计入量。原始阻塞区间始终保存在明细中，没有用缩短后的节点时长覆盖。')
    paragraph(doc,'表4是额外分类统计，不替代原始明细。前三个时间列为三轮累计ms，最后一列除以384场，逐项相加对应表2“全程调度与同步”。')
    rows=[[r[0],str(r[1]),f'{r[2]:.6f}',f'{r[3]:.6f}',f'{r[4]:.6f}',f'{r[5]:.6f}'] for r in raw['global_rows']]
    total=['合计',str(raw['original_control_records']),*[f'{math.fsum(r[j] for r in raw["global_rows"]):.6f}' for j in range(2,6)]]
    table(doc,['调度类别','条数','原始区间\n累计/ms','依赖分离后\n节点累计/ms','并行后\n计入/ms','等效时间\nms/16视频'],rows+[total],[1.55,.5,1.35,1.4,1.15,1.15],(1,2,3,4,5))
    paragraph(doc,f'各类别等效时间之和为{raw["global_ms16"]:.15f} ms，与表2及Excel附录中的全程调度行一致。它们分布在整轮进入、场间、槽位复用、真实捕获周边以及最后返回，不是Decoder之后额外串行执行的一段。')
    paragraph(doc,'对应三轮原始运行：'+ '、'.join(r['run'] for r in raw['sources'])+'。Excel原始分块表保留原始计时、依赖节点及批准时间账本的SHA256，逐条节点ID可定位原记录。所有保存值均为固定数值，不依赖其他工作表。')
    doc.save(target)
    print(json.dumps(dict(status='edited',path=str(target),source_sha256=digest(source),equations=5,
        original_control_records=raw['original_control_records'],preserved_original_document=True),ensure_ascii=True))


def audit(folder,pages,visual_reviewed):
    folder=Path(folder).resolve();target=folder/NAME;data=json.loads((folder/'static_values.json').read_text());p=data['appendix_table']
    d=Document(target);ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math','w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    with ZipFile(target) as z:
        xml=ET.fromstring(z.read('word/document.xml'))
        if len(xml.findall('.//m:oMath',ns))!=5:raise ValueError('Five native editable equations required')
    if len(d.tables)!=(4 if data.get('edit_mode')=='raw_appendix' else 2):raise ValueError('Unexpected document tables')
    timing=d.tables[1]
    if len(timing.rows)!=19:raise ValueError('17 accounted groups plus total and header required')
    for row,expected in zip(timing.rows[1:],p['rows']):
        values=[c.text for c in row.cells]
        if values!=[expected[0],expected[1],f'{expected[2]:.6f}']:raise ValueError('Word/Excel timing mismatch')
    text='\n'.join(x.text for x in d.paragraphs)+'\n'+'\n'.join(c.text for t in d.tables for r in t.rows for c in r.cells)
    if any(s in text for s in ['语言光场Adapter','文本Adapter”','[待补','TODO']):raise ValueError('Unresolved terminology or placeholder')
    if '不是物理光电整机的直接端到端实测' not in text:raise ValueError('Retain timing basis qualification')
    preserved_native_content=None
    if data.get('edit_mode') in ('optical_control_to_electronics','raw_appendix'):
        source=Path(data['base_word_document'])
        if digest(source)!=data['base_word_sha256']:raise ValueError('Source Word changed')
        with ZipFile(source) as old,ZipFile(target) as new:
            oldxml=ET.fromstring(old.read('word/document.xml'))
            if [ET.tostring(x) for x in oldxml.findall('.//m:oMath',ns)]!=[ET.tostring(x) for x in xml.findall('.//m:oMath',ns)]:raise ValueError('Native equations changed')
            names=sorted(n for n in old.namelist() if n.startswith('word/media/'))
            if names!=sorted(n for n in new.namelist() if n.startswith('word/media/')) or any(old.read(n)!=new.read(n) for n in names):raise ValueError('Existing illustration changed')
        if '统一合入紧接其后的同名OEO电子处理行' not in text:raise ValueError('Missing optical control scope')
        preserved_native_content=True
    if data.get('edit_mode')=='raw_appendix':
        raw=data['raw_appendix']
        expected=[[r[0]+' '+r[1].replace('\n',' '),*[f'{r[j]/384:.6f}' for j in [2,3,4]],f'{r[6]:.6f}'] for r in raw['summary_rows']]
        if [[c.text for c in r.cells] for r in d.tables[2].rows[1:]]!=expected:raise ValueError('Word raw summary mismatch')
        expected=[[r[0],str(r[1]),*[f'{r[j]:.6f}' for j in range(2,6)]] for r in raw['global_rows']]
        expected+=[['合计',str(raw['original_control_records']),*[f'{math.fsum(r[j] for r in raw["global_rows"]):.6f}' for j in range(2,6)]]]
        if [[c.text for c in r.cells] for r in d.tables[3].rows[1:]]!=expected:raise ValueError('Word original-control aggregate mismatch')
    images=sorted((folder/'word_native_render_final').glob('page-*.png'))
    if len(images)!=pages or pages<4 or not visual_reviewed:raise ValueError('Review every rendered page before sealing')
    result=dict(status='passed',native_Word_equations=5,word_Excel_timing_cells_identical=True,
        all_pages_visually_reviewed=True,pages=pages,document_sha256=digest(target),
        source_values_sha256=digest(folder/'static_values.json'),render_engine='Microsoft Word native PDF export then bundled Poppler',
        original_equations_and_illustration_preserved=preserved_native_content,
        stage_sum_ms=math.fsum(r[2] for r in p['rows'][:-1]),canonical_total_ms=data['rows'][0][5])
    with (folder/'document_audit.json').open('x',encoding='utf8') as f:json.dump(result,f,ensure_ascii=True,indent=2)
    print(json.dumps(result,ensure_ascii=True))


if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('mode',choices=['build','audit']);ap.add_argument('folder',type=Path)
    ap.add_argument('--revise',action='store_true');ap.add_argument('--pages',type=int);ap.add_argument('--visual-reviewed',action='store_true');a=ap.parse_args()
    build(a.folder,a.revise) if a.mode=='build' else audit(a.folder,a.pages,a.visual_reviewed)
