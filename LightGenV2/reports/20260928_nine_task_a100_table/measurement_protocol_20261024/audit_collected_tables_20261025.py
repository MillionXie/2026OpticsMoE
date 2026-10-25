"""Independent file-only audit of raw intervals/power and exported XLSX values."""
from pathlib import Path
from bisect import bisect_right
import csv,hashlib,json,math,sys,zipfile
from xml.etree import ElementTree as ET
from pypdf import PdfReader
from openpyxl import load_workbook
HERE=Path(__file__).resolve().parent;REPO=HERE.parents[3]
OUT=HERE/'outputs/01a061e2-6f88-7913-87f5-b632000e4c83/collected_20261025_01'
data=json.loads((OUT/'collected_data_and_audit.json').read_text(encoding='utf8'))
csv.field_size_limit(10000000)
def rows(p):
    with p.open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
def near(a,b,label):
    if not isinstance(a,(int,float)) or not math.isclose(a,b,rel_tol=1e-9,abs_tol=1e-7):raise ValueError(f'{label}: {a} != {b}')
def digest(p):return hashlib.sha256(p.read_bytes()).hexdigest()
seen={};raw=[]
for r in data['mainRuns']+data['tempRuns']:
    if r['folder'] in seen:continue
    seen[r['folder']]=r
    root=REPO/r['folder'];calls=rows(root/('calls.csv' if (root/'calls.csv').exists() else 'samples.csv'))
    assert len(calls)==200
    start_key='input_start' if 'input_start' in calls[0] else 'preprocess_start'
    end_key='cpu_result_return' if 'cpu_result_return' in calls[0] else 'call_return'
    starts=[float(x[start_key]) for x in calls];ends=[float(x[end_key]) for x in calls]
    assert all(a<=b for a,b in zip(starts,ends))
    gaps=[starts[i]-ends[i-1] for i in range(1,200)];assert min(gaps)>=-1e-8
    if 'internal_gap_ms' in calls[0]:
        for i,x in enumerate(calls):near(float(x['internal_gap_ms']),0 if i==0 else gaps[i-1]*1000,'all199 original gaps')
    entry={'folder':r['folder'],'calls':len(calls),'internal_gaps':199,'internal_gap_seconds':sum(gaps)}
    if r['system']=='ours':
        stages=rows(root/'stages.csv');chosen=[x for x in stages if x['charged']=='True']
        for x in stages:near((float(x['end_monotonic'])-float(x['start_monotonic']))*1000,float(x['synchronized_wall_ms']),'stage sync wall')
        charged=sum(float(x['synchronized_wall_ms']) for x in chosen)/1000+sum(gaps)
        near(charged,r['charged_s'],'original charged stage union +199gaps')
        near(charged+r['optical_s'],r['seconds200'],'hybrid duration')
        near(104.230*charged+121.746*r['optical_s'],r['joules200'],'phase J')
        entry['charged_stage_count']=len(chosen)
    else:
        start,end=starts[0],ends[-1];near(end-start,r['seconds200'],'whole baseline wall')
        power=rows(root/'telemetry.csv');ts=[float(x['host_monotonic_s']) for x in power]
        assert ts[0]<=start and ts[-1]>=end and all(b>a for a,b in zip(ts,ts[1:]))
        energy=0.;age=0.;cursor=start
        while cursor<end:
            i=bisect_right(ts,cursor)-1;right=min(end,ts[i+1]);assert right>cursor
            age=max(age,right-ts[i]);energy+=float(power[i]['watts'])*(right-cursor);cursor=right
        assert age<=.25000001
        near(338.2*(end-start)+energy,r['joules200'],'independent same-window baseline energy')
        entry.update(board_energy_j=energy,max_support_age_s=age)
    raw.append(entry)
assert len(raw)==78

workbooks=[]
for filename,records,runs,main in [
 ('A100_九行任务_已采集核验_20261025.xlsx',data['mainRecords'],data['mainRuns'],True),
 ('A100_Temporal多基线_已采集核验_20261025.xlsx',data['tempRecords'],data['tempRuns'],False)]:
    p=OUT/filename;wb=load_workbook(p,data_only=True,read_only=True);formulas=load_workbook(p,data_only=False,read_only=True)
    errors=[]
    for s in wb:
        for row in s:
            for c in row:
                if c.data_type=='e':errors.append([s.title,c.coordinate,c.value])
    assert not errors,errors
    s=wb['指标与计算']
    for row,r in enumerate(records,6):
        assert s.cell(row,1).value==r['key']
        if not r.get('calls'):
            assert all(s.cell(row,c).value=='待测' for c in range(11,17));continue
        own=[x for x in runs if x['key']==r['key']];seconds=sum(x['seconds200'] for x in own);joules=sum(x['joules200'] for x in own)
        for c,v in [(11,seconds),(12,joules),(13,seconds/600*1000),(14,joules/600),(15,600/seconds),(16,r['tokens']/seconds)]:near(s.cell(row,c).value,v,'exported cached aggregate')
        if r.get('op'):
            for c,v in [(17,r['op']['OP600']/600/1e12),(18,r['op']['OP600']/seconds/1e12),(19,r['op']['OP600']/joules/1e12)]:near(s.cell(row,c).value,v,'exported cached NN')
    if main:
        summary=wb['九行对比']
        for i in range(9):
            o,q=records[2*i:2*i+2];row=i+6
            if o.get('calls') and q.get('calls'):
                speed=q['ms']/o['ms'];energy=q['j']/o['j']
                for c,v in [(8,speed),(11,energy),(12,1-1/energy),(13,speed*energy),(14,speed)]:near(summary.cell(row,c).value,v,'main ratio cache')
            else:assert summary.cell(row,8).value=='待测'
    else:
        summary=wb['Temporal对比']
        for row,r in enumerate(records,6):
            if not r.get('calls'):continue
            anchor=next(x for x in records if x['system']=='qwen' and x['cap']==100 and x['power_limit']==r['power_limit'])
            for c,v in [(9,anchor['ms']/r['ms']),(10,anchor['j']/r['j']),(11,1-r['j']/anchor['j']),(12,anchor['ms']/r['ms']*anchor['j']/r['j'])]:near(summary.cell(row,c).value,v,'same-platform anchor ratio')
    fcount=sum(c.data_type=='f' for s in formulas for row in s for c in row)
    assert fcount>200
    workbooks.append({'file':filename,'bytes':p.stat().st_size,'sha256':digest(p),'formulas':fcount,'formula_errors':0,'cached_values_independently_verified':True})
    wb.close();formulas.close()
result={'status':'passed','no_models_no_GPU':True,'independent_raw_points':78,'original_calls':15600,'original_internal_gaps':15522,'raw_audits':raw,'workbooks':workbooks,'scope':'File-only original interval/power recomputation plus exported XLSX cached values; does not rerun full numerical/model contracts'}
(OUT/'independent_delivery_audit.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps({k:v for k,v in result.items() if k not in ('raw_audits',)},ensure_ascii=False))

if '--visual-review-accepted' in sys.argv:
    doc=OUT/'A100_时间能耗与倍数计算说明_20261025.docx'
    with zipfile.ZipFile(doc) as z:xml=ET.fromstring(z.read('word/document.xml'))
    equations=len(xml.findall('.//{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath'))
    pdf=OUT/'docx_qa/methods_native_word.pdf'
    pages=len(PdfReader(pdf).pages)
    assert equations==20 and pages==6
    names=[x['file'] for x in workbooks]+[doc.name,'raw_file_index.csv','independent_delivery_audit.json','collected_data_and_audit.json']
    manifest={
      'status':'PARTIAL_COLLECTED_DATA_NOT_FINAL_ALL18',
      'main_cost_cells_complete':14,'main_cost_cells_total':18,'main_complete_pairs':6,
      'temporal_baseline_trials_complete':36,'supplementary_80GB_Qwen_trials_complete':3,
      'source_indices':[data['main_index'],data['temporal_index']],
      'raw_file_catalog':{'total':915,'rechecked_against_original_sealed_SHA_size':849,'auxiliary_manifest_receipt_local_audit_with_new_SHA':66},
      'review':{'independent_original_raw_points':78,'xlsx_formula_errors':0,'visually_checked_worksheet_previews':11,'Word_native_editable_equations':equations,'Word_rendered_pages_visually_checked':pages},
      'outputs':[{'file':name,'path':str(OUT/name),'bytes':(OUT/name).stat().st_size,'sha256':digest(OUT/name)} for name in names],
      'limitations':['Four main cost cells remain pending; no historical workbook backfill','Ours hybrid hardware projection is not whole-software wall or physically measured device E2E','Baseline energy includes fixed host proxy, not wall-meter measurement','Physical optical dense-equivalent OP remains unknown; NN subtotal is not all scalar operations'],
      'no_models_no_GPU_no_rental_connection':True}
    target=OUT/'delivery_manifest.json'
    target.write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf8')
    print(json.dumps({'manifest':str(target),'sha256':digest(target),'docx':manifest['outputs'][2]},ensure_ascii=False))
