"""Independently read original CSV intervals and exported timing XLSX caches."""
from pathlib import Path
from collections import defaultdict
import csv,hashlib,json,math
from openpyxl import load_workbook
HERE=Path(__file__).resolve().parent;REPO=HERE.parents[3]
ROOT=HERE/'outputs/01a061e2-6f88-7913-87f5-b632000e4c83'
OUT=ROOT/'optical_breakdown_20261025_01'
data=json.loads((ROOT/'collected_20261025_01/collected_data_and_audit.json').read_text(encoding='utf8'))
audit=json.loads((OUT/'timing_breakdown_audit.json').read_text(encoding='utf8'))
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def near(a,b):assert isinstance(a,(int,float)) and math.isclose(a,b,rel_tol=1e-9,abs_tol=1e-6),(a,b)
def key(x):return f'{int(x):02}' if isinstance(x,(int,float)) else x
def rows(p):
    with p.open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
group=defaultdict(lambda:[0.,0]);windows=defaultdict(float)
for r in data['mainRuns']:
    if r['system']!='ours':continue
    k=r['key'].split('|')[0];root=REPO/r['folder'];calls=rows(root/('calls.csv' if (root/'calls.csv').exists() else 'samples.csv'));stages=rows(root/'stages.csv')
    assert len(calls)==200
    a='input_start' if 'input_start' in calls[0] else 'preprocess_start';b='cpu_result_return' if 'cpu_result_return' in calls[0] else 'call_return'
    start=float(calls[0][a]);end=float(calls[-1][b]);software=(end-start)*1000;windows[k]+=software
    gap=sum(float(calls[i][a])-float(calls[i-1][b]) for i in range(1,200))*1000
    intervals=sorted((float(s['start_monotonic']),float(s['end_monotonic'])) for s in stages)
    assert all(start-1e-8<=l<=u<=end+1e-8 for l,u in intervals)
    assert all(y<=z+1e-8 for (x,y),(z,t) in zip(intervals,intervals[1:]))
    total=charged=0.
    for s in stages:
        ms=(float(s['end_monotonic'])-float(s['start_monotonic']))*1000;near(ms,float(s['synchronized_wall_ms']));total+=ms
        flag=s['charged']=='True';charged+=ms if flag else 0
        g=group[k,s['stage'],s['family'],flag];g[0]+=ms;g[1]+=1
    near(charged+gap,r['charged_s']*1000)
    for stage,family,flag,ms,count in [('__inter_request_gap','request_gap',True,gap,199),('__unspanned_software_overhead','software_remainder',False,software-total-gap,200)]:
        g=group[k,stage,family,flag];g[0]+=ms;g[1]+=count
assert len(group)==audit['stage_rows']
file=OUT/audit['output']['file'];wb=load_workbook(file,data_only=True,read_only=True);form=load_workbook(file,data_only=False,read_only=True)
errors=[(s.title,c.coordinate,c.value) for s in wb for row in s for c in row if c.data_type=='e'];assert not errors
s=wb['任务时间拆分'];records=[x for x in data['mainRecords'] if x['system']=='ours']
for n,r in enumerate(records,6):
    k=r['key'].split('|')[0];assert key(s.cell(n,1).value)==k
    cycles=(13/200 if k=='01a' else 1)*(3 if k in ('04','05','06') else 6)
    expected=[cycles,cycles*.714,cycles*.3,cycles*.0307,cycles*1.0447]
    for c,v in enumerate(expected,3):near(s.cell(n,c).value,v)
    if r.get('calls'):
        electronic=sum(v[0] for g,v in group.items() if g[0]==k and g[3])/600
        near(s.cell(n,8).value,electronic);near(s.cell(n,9).value,r['ms']);near(electronic+expected[-1],r['ms']);near(s.cell(n,10).value,expected[-1]/r['ms'])
    else:assert all(s.cell(n,c).value=='待测' for c in (8,9,10))
d=wb['电子与仿真明细'];seen=set()
for row in range(6,6+len(group)):
    item=(key(d.cell(row,1).value),d.cell(row,4).value,d.cell(row,6).value,d.cell(row,5).value=='是');assert item in group and item not in seen;seen.add(item)
    near(d.cell(row,7).value,group[item][0]);near(d.cell(row,8).value,group[item][1]);near(d.cell(row,9).value,group[item][0]/600)
assert seen==set(group)
o=wb['光路逐步骤']
for row in range(6,6+audit['optical_step_rows']):
    k=key(o.cell(row,1).value);cycles=13/200 if k=='01a' else 1
    for c,v in [(7,13 if k=='01a' else 200),(8,cycles*.714),(9,cycles*.3),(10,cycles*.0307),(11,cycles*1.0447)]:near(o.cell(row,c).value,v)
done=[x for x in records if x.get('calls')]
families=[['input'],['frontend'],['bridge'],['adapter','amplitude_encoding_nn'],['router'],['ccd_nonlinearity','ccd_readout'],['head','search','output'],['request_gap']]
for i,r in enumerate(done):
    k=r['key'].split('|')[0];near(s.cell(46+i,3).value,windows[k]/600)
    for c,fs in enumerate(families,3):near(s.cell(67+i,c).value,sum(v[0] for g,v in group.items() if g[0]==k and g[2] in fs)/600)
    near(s.cell(46+i,8).value,s.cell(67+i,11).value)
for ref in audit['source_files']:
    p=REPO/ref['path'];assert p.stat().st_size==ref['bytes'] and sha(p)==ref['sha256']
result={'status':'passed','scope':'Original interval arithmetic and saved XLSX cached values only; no model or GPU work','measured_Ours_runs':18,'measured_valid_calls':3600,'internal_request_gaps':3582,'original_stage_groups_checked':len(group),'optical_steps_checked':audit['optical_step_rows'],'formula_errors':errors,'cached_values_verified':True,'source_files_verified':len(audit['source_files']),'output':audit['output']}
(OUT/'independent_timing_breakdown_audit.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(result,ensure_ascii=False))
