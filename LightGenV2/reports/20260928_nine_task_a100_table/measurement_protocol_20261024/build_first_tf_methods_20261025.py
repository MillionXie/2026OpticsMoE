"""Chinese editable-equation appendix for the first-compute-module delivery."""
from pathlib import Path
from copy import deepcopy
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
HERE=Path(__file__).resolve().parent
OUT=HERE/'outputs/01a061e2-6f88-7913-87f5-b632000e4c83/first_tf_complete_20261025_01'
assert not (OUT/'delivery_manifest.json').exists()
data=json.loads((OUT/'core_data.json').read_text(encoding='utf8'));assert not any(x['pending'] for x in data['mainRecords'])
D=Document();section=D.sections[0];section.page_width=Inches(8.5);section.page_height=Inches(11)
section.left_margin=section.right_margin=Inches(.72);section.top_margin=section.bottom_margin=Inches(.65)
for name in ('Normal','Title','Heading 1','Heading 2'):
    s=D.styles[name];s.font.name='Times New Roman';s.font.color.rgb=RGBColor(0,0,0)
    s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
    pp=s.element.find(qn('w:pPr'))
    if pp is not None:
        for e in pp.findall(qn('w:pBdr')):pp.remove(e)
D.styles['Normal'].font.size=Pt(11);D.styles['Normal'].paragraph_format.line_spacing=1.08
D.styles['Normal'].paragraph_format.space_after=Pt(6)
D.styles['Title'].font.size=Pt(20);D.styles['Heading 1'].font.size=Pt(14)
D.core_properties.title='A100首计算模块时间能量和倍数计算说明'
D.core_properties.author='LightGen2 Research'
def p(t='',style=None):return D.add_paragraph(t,style)
def h(t):return p(t,'Heading 1')
def page(t):D.add_page_break();h(t)
def text(t):
    r=OxmlElement('m:r');z=OxmlElement('m:t');z.text=str(t);r.append(z);return r
def el(tag,parts):
    z=OxmlElement(tag)
    for x in parts:z.append(text(x) if isinstance(x,str) else deepcopy(x))
    return z
def sub(a,b):return el('m:sSub',[el('m:e',[a]),el('m:sub',[b])])
def frac(a,b):return el('m:f',[el('m:num',a),el('m:den',b)])
def total(symbol,lo,hi,body):
    z=OxmlElement('m:nary');pr=OxmlElement('m:naryPr');ch=OxmlElement('m:chr');ch.set(qn('m:val'),symbol);pr.append(ch);z.append(pr)
    z.append(el('m:sub',[lo]));z.append(el('m:sup',[hi]));z.append(el('m:e',body));return z
count=0
def eq(*parts):
    global count
    count+=1;z=p();z.alignment=WD_ALIGN_PARAGRAPH.CENTER;z.paragraph_format.keep_together=True
    z._p.append(el('m:oMathPara',[el('m:oMath',[*parts,'   ('+str(count)+')'])]))
def table(headers,rows,widths):
    t=D.add_table(rows=1,cols=len(headers));t.autofit=False
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        z=OxmlElement('w:'+edge);z.set(qn('w:val'),'single');z.set(qn('w:sz'),'4');z.set(qn('w:color'),'D9D9D9');borders.append(z)
    t._tbl.tblPr.append(borders)
    for i,values in enumerate([headers]+rows):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        for j,(c,value) in enumerate(zip(cells,values)):
            c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;c.text=str(value)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'243A50' if i==0 else 'F1F5F8' if i%2==0 else 'FFFFFF');c._tc.get_or_add_tcPr().append(shade)
            for para in c.paragraphs:
                para.paragraph_format.space_before=Pt(4);para.paragraph_format.space_after=Pt(4)
                if j>0:para.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for r in para.runs:
                    r.font.size=Pt(10);r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
                    if i==0:r.bold=True;r.font.color.rgb=RGBColor(255,255,255)
        if i==0:t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    p();return t
records={r['key']:r for r in data['mainRecords']};lg=records['01a|ours'];q=records['01a|qwen']
p('A100首计算模块时间能量和倍数计算说明','Title')
p('九任务主表与Temporal多基线表的方法附录 2026年10月25日')
p('本版比较从模型主体的首个计算模块开始的时间与能量。Qwen从原生首个Transformer块进入，LG从首个光学模块的输入映射进入，计到原任务结果返回CPU。进入起点之前的文件读取、解码、分词、词向量和图像位置准备不计入主结果。')
p(f'Temporal的LG按16视频、64帧整幅相干光场实际执行。折16个有效视频的同步耗时为{lg["ms"]*16:.4f} ms，Qwen100为{q["ms"]*16:.4f} ms，时间比为{q["ms"]/lg["ms"]:.3f}倍。等效单帧分别再除64，倍数不变。该结果是核心范围比较，不是从视频文件开始的全流程比较。')
table(['任务','Qwen除LG时间比','Qwen除LG能耗比','节能率'],[
    [records[k+'|ours']['task'],f'{records[k+"|qwen"]["ms"]/records[k+"|ours"]["ms"]:.3f}',
     f'{records[k+"|qwen"]["j"]/records[k+"|ours"]["j"]:.3f}',
     f'{1-records[k+"|ours"]["j"]/records[k+"|qwen"]["j"]:.2%}']
    for k in ('01a','01b','02','03','04','05','06','07','08')],[2.5,1.45,1.6,1.3])
p('正式精度沿用原模型完整TEST。三轮各200次仅用于时间、能量和输出核验，不用首200条的子集指标替代完整TEST。80GB与40GB对照单独标明平台，不能把跨平台差异归因为纯模型收益。')

page('1 计时起点和保留的计算')
p('图像类Qwen从原生视觉Transformer第0块进入，纯文字与T12从原生语言Transformer第0块进入。Temporal基线使用原封存的首块时间戳，其余七项Qwen按本次授权补记相同位置。YOLO没有Transformer，沿用原生第一个卷积块起点，单独标注。')
p('LG按原记录中第一个光学模块的输入映射切分，不从更晚的光传播开始。光模块内部幅度映射、路由、数字读出，后续真实合并器、任务输出头、图像生成网络、检索排序和结果返回均按原计费规则保留。')
p('“不加前置”指不把进入首计算模块之前的时间加回去，也不使用三者共用embedding耗时的代数情景。起点之后若原模型仍执行另一个分支的词向量或图像前端，这些后续计算保留，不能借切分口径删掉原网络的中间步骤。')
table(['项目','含义','本版处理'],[
 ['CPU输入准备','处理器读取文件，解码、抽帧、裁剪缩放并整理输入','进入起点前不计'],
 ['内存到显存传输','把CPU内存中的输入复制到GPU显存，旧表缩写H2D','进入起点前不计'],
 ['文字分词 tokenizer','CPU把文字变成词编号和有效位置标记，不是词向量查表','进入起点前不计'],
 ['词向量 embedding','GPU按词编号查表得到数值特征。图像分块嵌入是卷积或投影','进入起点前不计'],
 ['位置处理','位置特征构造、插值或旋转位置编码','进入起点前不计'],
 ['主体和任务头','原完整主网络，以及生成评分、坐标、图片或检索结果的输出网络','起点之后保留'],
 ['CPU输出','把最终结果送回处理器内存，执行原输出处理','保留'],
 ['请求间隙','上一次结果返回到下一次开始准备输入之间的实际空档','本版核心区间不计']],[1.6,3.5,1.75])

page('2 三轮汇总和并行单位换算')
p('每个补测基线使用三个独立新进程，每进程顺序执行原首200个真实输入，无dummy或专门推理预热，保留首条与全部慢调用。原模型、精度、输入和排序规则不变，完整输出与原合格结果核对。额外观察只记录模块入口及结果传输事件，不增加模型前向。')
eq(sub('T','Q,r'),' = ',total('∑','i=1','200',[sub('t','CPU结果,r,i'),' − ',sub('t','首TF,r,i')]))
p('每次调用单独取首块入口到CPU任务结果的区间。只对这200个区间求和，不把区间之间的输入准备与间隙混入核心总时间。')
eq(sub('T','LG,r'),' = ',sub('T','选定电子,r'),' + ',sub('T','固定光,r'))
p('LG电子时间采用起点之后原选定互斥阶段的同步墙钟之和。真实运行的GPU光传播仿真、布局、融合和原批准单列的并行电子补偿不加入硬件投影。连续软件运行时间保留为诊断，不能冒称物理装置端到端实测。')
eq(sub('t','单光周期'),' = 0.714 + 0.300 + 0.0307 = 1.0447 ms')
p('三个固定分量分别为空间光调制器更新、曝光和额外相机读出。一般输入6周期为6.2682 ms，LSP、SALICON和纯标题在线查询为3周期3.1341 ms。这些是既有固定参数，不是本轮新测光路。')
eq(sub('T','16视频'),' = ',frac([sub('T','三轮合计')],['600']),' × 16')
eq(sub('T','等效单帧'),' = ',frac([sub('T','16视频')],['64']))
p('Temporal每轮13场、每场16视频4帧，共208执行槽，200个有效视频。三轮39场执行624槽、600个有效视频，末场8个补齐槽的成本完整保留。因此折16有效视频的固定光时间为6.518928 ms，而一次满场的固定光时间为6.2682 ms。Spatial保持原模型每场1视频4帧，不将现有时间再除16。')
p('表内显卡事件耗时是GPU事件观察区间，用来核对GPU执行，不与同步墙钟相加。主表倍率和能量时延积均使用同步耗时，避免将GPU事件分母与另一范围的能量混算。')

page('3 能量和比较倍数')
p('所有能量公式中的时间使用秒。LG能量按既有阶段功率模型估算，基线使用338.2 W主机功率代理加相同核心区间的实际板卡功率积分。二者都不能称为插座电表实测。')
eq(sub('E','LG'),' = 104.230 ',sub('T','选定电子'),' + 121.746 ',sub('T','固定光'))
eq(sub('E','Q'),' = 338.2 ',sub('T','Q'),' + ',total('∑','r,i','',[total('∫',sub('t','首TF,r,i'),sub('t','CPU结果,r,i'),[sub('P','板卡'),'(t) dt'])]))
p('基线重新积分每个首块至CPU结果的真实区间，不把旧完整输入窗口的能量按时间比例裁减。功率观测采用原零阶保持规则，区间前后均有支持数据，年龄上限250 ms。不扣idle，不加入区间外输入准备或间隙。LG原投影排除的传播仿真和并行补偿能量不加回。')
eq(sub('S','时间'),' = ',frac([sub('T','Q')],[sub('T','LG')]))
eq(sub('R','能耗'),' = ',frac([sub('E','Q')],[sub('E','LG')]))
eq(sub('η','节能'),' = 1 − ',frac([sub('E','LG')],[sub('E','Q')]))
eq(sub('R','能量时延积'),' = ',frac([sub('E','Q'),sub('T','Q')],[sub('E','LG'),sub('T','LG')]),' = ',sub('S','时间'),sub('R','能耗'))
p('时间比或能耗比大于1表示LG更快或每输入能量更低。等效单帧的两边同时除4或64，比例不变。共同词编号吞吐使用同一批Qwen有效非padding词编号数作为双方分子，表示同任务归一化速率，不是生成语言的token吞吐。')
eq(sub('V','共同词编号'),' = ',frac([sub('N','同批有效词编号')],[sub('T','核心')]))
p('新范围的完整运算量没有重新独立核定。原完整网络的NN dense-MAC不能直接配新核心时间计算TOPS或TOPS/W。若后续核定，1个实数MAC按2 OP；物理光dense-equivalent OP未知，不能填0，不能用容量或FFT次数代替。')
p('各Excel保留三轮原值、起点名称和原文件路径，LG逐阶段给出中文含义。原数据索引登记精确路径、文件大小及SHA256；本版生成后独立重算区间能量与表格公式缓存，并检查文档每一页。旧全流程版与各情景表原样留档。')
D.save(OUT/'A100_首计算模块时间能量与倍数说明_20261025.docx')
print(json.dumps(dict(equations=count,output=str(OUT/'A100_首计算模块时间能量与倍数说明_20261025.docx'))))
