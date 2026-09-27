from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
DATA = json.loads((ROOT / "calculated_summary.json").read_text(encoding="utf-8"))


def shade(cell, color: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    fill = OxmlElement("w:shd")
    fill.set(qn("w:fill"), color)
    props.append(fill)


def set_cell_text(cell, value, bold=False, color=None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(str(value))
    run.bold = bold
    run.font.size = Pt(8)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(1.8); sec.bottom_margin = Cm(1.8)
sec.left_margin = Cm(1.8); sec.right_margin = Cm(1.8)
styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
styles["Normal"].font.size = Pt(9)
styles["Title"].font.name = "Arial"
styles["Title"]._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
styles["Title"].font.size = Pt(18)
styles["Heading 1"].font.name = "Arial"
styles["Heading 1"]._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
styles["Heading 1"].font.size = Pt(12)

title = doc.add_heading("A100 光电 MoE 与 Qwen3VL 基线：时间、功率与能效测量说明", 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
p = doc.add_paragraph("固定版本 · 200 次正式测量 · 无专门预热 · 第 1 次计入")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.runs[0].italic = True

doc.add_heading("摘要", level=1)
doc.add_paragraph(
    "本文件定义并复核七项已完成任务的推理边界、计时方法、功率输入与能量公式。"
    "04（ABO 文搜图）与 08（ABO clean 文生图）仅预留，不填入任何推测结果。性能取绑定检查点的完整测试集结果；"
    "基线速度取测试集顺序前 200 项；Ours 对固定形状、GPU 驻留的任务电子核连续执行 200 次。"
    "两者均仅加载模型一次，不执行显式预热，且保留第 1 次测量。"
)

doc.add_heading("测量边界与时钟", level=1)
doc.add_paragraph(
    "Qwen3VL 基线使用 Synchronized Wall：在第一个原生 Vision Transformer block 的 pre-hook 处开始，"
    "经过所需 Vision/Language blocks 与任务读出头，在最终输出就绪并完成 CUDA synchronize 后停止。"
    "模型加载、文件解码、processor/tokenizer 与 CPU→GPU 搬运不在该区间内。"
)
doc.add_paragraph(
    "Ours 的电子时间使用 CUDA Event，并仅串行计入 CCD 后非线性与学习式读出、光路由算术、必要 bridge/merger 与任务头；"
    "与光传播并行的残差支路不进入关键路径。每个电子核执行 200 次，报告全部 200 次（含第 1 次）的均值。"
    "CCD 裁剪/堆叠、排版、SLM 编码、文件 I/O、搬运以及通用整场张量融合不进入该窄口径。"
)
doc.add_paragraph(
    "单次物理传播固定为 1.0447 ms = 0.714 ms（高速 SLM）+ 0.300 ms（曝光）+ 0.0307 ms（812×812 ROI 相机读出）。"
    "含语言的任务与双层视频质量网络为 6 次传播；LSP 与 SALICON 为仅视觉的 3 次传播。"
)

doc.add_heading("功率与能量计算", level=1)
doc.add_paragraph(
    "光学设备功率为 80.358 W：激光器 9.679 W、相位 SLM 5.595 W、CCD 14.303 W、"
    "高速 SLM 34.225 W + 16.556 W。控制机箱固定 41.388 W；Ours 的 A100 固定 62.842 W。"
    "基线机箱固定 338.2 W，A100 板卡功率按各任务连续推理阶段的 nvidia-smi active mean 读取，不做空载扣除。"
)
for formula in (
    "Ours：E = t光 × (80.358 + 41.388) + t电 × (41.388 + 62.842)",
    "Baseline：E = tWall × (338.2 + P_A100,active)",
    "速度提升 = tBaseline / tOurs；能效提升 = EBaseline / EOurs；能量降低 = 1 − EOurs / EBaseline。",
):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    r = p.add_run(formula); r.bold = True

doc.add_heading("主要结果", level=1)
headers = ["序号", "任务", "Ours ms", "Baseline ms", "Ours J", "Baseline J", "速度×", "能效×"]
table = doc.add_table(rows=1, cols=len(headers))
table.style = "Table Grid"
for i, h in enumerate(headers):
    set_cell_text(table.rows[0].cells[i], h, True, "FFFFFF")
    shade(table.rows[0].cells[i], "0B5A88")
for i, row in enumerate(DATA["rows"]):
    vals = [row["number"], row["task"], f'{row["ours_total_time_ms"]:.3f}', f'{row["baseline_total_time_ms_matched_workload"]:.3f}', f'{row["ours_total_energy_j"]:.3f}', f'{row["baseline_total_energy_j_matched_workload"]:.3f}', f'{row["speedup_x"]:.2f}', f'{row["energy_efficiency_gain_x"]:.2f}']
    cells = table.add_row().cells
    for j, value in enumerate(vals):
        set_cell_text(cells[j], value)
        shade(cells[j], "E7EAED" if i % 2 else "DCECF7")

doc.add_paragraph(
    "注：LGVQ temporal 的 Ours 是 16 个视频共享一个光场的批次延迟；基线为 16 次 batch=1 调用之和。"
    "其余行均为单样本。表中所有派生数值由同目录 calculated_summary.json/csv 生成。"
)

doc.add_heading("运算量口径", level=1)
doc.add_paragraph(
    "统一采用 1 个实数 MAC = 2 OP。基线通过真实前向钩子统计第一个 Vision block 后实际执行的 Linear、Conv 与注意力 QK/AV；"
    "不把激活、归一化、池化、索引、softmax 等无 MAC 操作虚构为 MAC。Ours 的传播运算量按 478×478 个模式的稠密复数线性变换等效计数，"
    "并明确标注为 effective TOP；A100 侧标注为 achieved MAC-derived TOP。二者不能与芯片数据手册峰值 TOPS 混用。"
)
ops_headers = ["序号", "Ours effective TOP", "Baseline achieved TOP", "Ours TOPS/W", "Baseline TOPS/W", "提升×"]
ops_table = doc.add_table(rows=1, cols=len(ops_headers))
ops_table.style = "Table Grid"
for i, h in enumerate(ops_headers):
    set_cell_text(ops_table.rows[0].cells[i], h, True, "FFFFFF")
    shade(ops_table.rows[0].cells[i], "0B5A88")
for i, row in enumerate(DATA["rows"]):
    vals = [row["number"], f'{row["ours_effective_top"]:.4f}', f'{row["baseline_achieved_top_matched_workload"]:.4f}', f'{row["ours_effective_tops_per_w"]:.4f}', f'{row["baseline_achieved_tops_per_w"]:.4f}', f'{row["tops_per_w_gain_x"]:.2f}']
    cells = ops_table.add_row().cells
    for j, value in enumerate(vals):
        set_cell_text(cells[j], value)
        shade(cells[j], "E7EAED" if i % 2 else "FCE4D6")

doc.add_heading("版本、数据与可复核性", level=1)
doc.add_paragraph(
    "每项报告均保留：逐样本/逐调用时间 CSV、连续功率采样 CSV、进程占用审计、模型/检查点/源码 SHA-256、运行命令与环境。"
    "原始服务器目录为 /DATA/DATA1/guest3/2026OpticsMoE_a100_measurements/20260927_demo_energy_efficiency_a100；"
    "本地镜像位于本说明同目录的 raw_remote，并另按 baseline/ours/任务编号整理。"
)
doc.add_paragraph(
    "LSP 的红框展示值 0.7353 尚未找到可严格绑定的结果文件；当前可追溯的 Ours 完整测试结果为 0.7347857143。"
    "为避免学术性替换，本报告保留精确绑定值。OpenMoji 基线固定为新布局 0.8120，不混用旧结果 0.8420。"
)

doc.add_heading("质量控制", level=1)
qc = [
    "A100 物理索引固定为 6；每个正式 profiler 在测量前后检查无其他计算 PID。",
    "01 的首轮功率误读物理 GPU 0 已被识别并作废；正式结果为重测后的物理 GPU 6 数据，作废文件保留在 audit_rejected_wrong_gpu0_power。",
    "所有计时均无显式 warm-up，且第 1 项包含在均值中；不通过删除慢启动样本缩短结果。",
    "04、08 状态为 reserved，待模型定稿后按相同协议补测。",
]
for item in qc:
    doc.add_paragraph(item, style="List Bullet")

footer = doc.sections[0].footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("LightGenV2 · A100 demo energy-efficiency audit · 2026-09-27").font.size = Pt(8)

out = ROOT / "A100_demo_energy_efficiency_methods.docx"
doc.save(out)
print(out)
